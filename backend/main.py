from fastapi import FastAPI, HTTPException, status, UploadFile, File, Form, Body, BackgroundTasks, Request, Depends
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from fastapi.responses import JSONResponse
from pydantic import BaseModel
import uuid
import json
import os
try:
    import redis
    redis_url = os.getenv('REDIS_URL', 'redis://localhost:6379/0')
    redis_client = redis.Redis.from_url(redis_url, decode_responses=True)
    redis_client.ping()
except Exception:
    redis = None
    redis_client = None

USE_CELERY = os.getenv('USE_CELERY', 'false').strip().lower() in {'true', '1', 'yes'}

class JobStore:
    def __init__(self):
        self.local_jobs = {}
        
    def __setitem__(self, key, value):
        if redis_client:
            redis_client.set(f"job:{key}", json.dumps(value), ex=86400) # expire in 24h
        else:
            self.local_jobs[key] = value
            
    def __getitem__(self, key):
        if redis_client:
            data = redis_client.get(f"job:{key}")
            if data:
                return json.loads(data)
            raise KeyError(key)
        return self.local_jobs[key]
        
    def __contains__(self, key):
        if redis_client:
            return redis_client.exists(f"job:{key}")
        return key in self.local_jobs

    def get(self, key, default=None):
        if key in self:
            return self[key]
        return default

GENERATION_JOBS = JobStore()

from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from typing import List, Optional
from datetime import datetime
import uvicorn
import os
import warnings
import shutil
import json
import sqlite3
import re
import tempfile
import time
import hashlib
import base64
import jwt

JWT_SECRET = os.getenv('JWT_SECRET', 'super-secret-key-change-me-to-something-secure-for-production')
JWT_ALGORITHM = 'HS256'
ALLOWED_ORIGINS = [
    origin.strip()
    for origin in os.getenv(
        'ALLOWED_ORIGINS',
        'http://localhost:5173,http://localhost:5174,http://localhost:4173,http://localhost:3000,http://127.0.0.1:5173,http://127.0.0.1:5174,http://127.0.0.1:4173,http://127.0.0.1:3000'
    ).split(',')
    if origin.strip()
]
from pathlib import Path

from email.message import EmailMessage

warnings.filterwarnings(
    "ignore",
    message=".*on_event is deprecated.*",
    category=DeprecationWarning,
)
warnings.filterwarnings("ignore", category=DeprecationWarning)

NUMBER_WORDS = {
    "one": 1,
    "two": 2,
    "three": 3,
    "four": 4,
    "five": 5,
    "six": 6,
    "seven": 7,
    "eight": 8,
    "nine": 9,
    "ten": 10,
}


def _effective_count_from_instruction(instruction: str, configured_count: int) -> int:
    text = (instruction or "").strip().lower()
    match = re.search(r"answer\s+any\s+(\d+|one|two|three|four|five|six|seven|eight|nine|ten)", text)
    if match:
        token = match.group(1)
        parsed = int(token) if token.isdigit() else NUMBER_WORDS.get(token)
        if parsed is not None:
            return max(0, min(parsed, configured_count))
    return max(0, configured_count)

# Core imports
from core.database import get_db_connection, init_database, get_db_type, get_cursor, get_placeholder
from core.models import (
    SubjectCreate, SubjectUpdate, SubjectResponse,
    QuestionBankCreate, QuestionBankResponse,
    QuestionCreate, QuestionResponse,
    BlueprintCreate, BlueprintResponse,
    QuestionPaperCreate, QuestionPaperResponse,
    QuestionGenerationRequest
)

# Utils imports
from utils.syllabus_parser import parse_syllabus, save_syllabus_to_db

# Services imports
from services.default_blueprint import DEFAULT_BLUEPRINT_STRUCTURE
from services.paper_generator import generate_question_paper
from services.qdrant_client import sync_subject_files_to_qdrant
from services.rag_retrieval import retrieve_context, format_context_for_prompt
from services.question_generator import generate_questions_with_ollama, test_ollama_connection
from services.grading_engine import generate_answer_script, grade_student_paper, extract_text_from_pdf

try:
    from core.celery_app import celery_app
except Exception:
    class DummyCelery:
        def task(self, *args, **kwargs):
            def decorator(f): return f
            return decorator
    celery_app = DummyCelery()

@celery_app.task(name="celery_run_generate_questions")
def celery_run_generate_questions(job_id: str, subject_id: int, request_dict: dict):
    request = QuestionGenerationRequest(**request_dict)
    run_generate_questions(job_id, subject_id, request)

@celery_app.task(name="celery_run_generate_all_questions")
def celery_run_generate_all_questions(job_id: str, subject_id: int, requests_dict: list):
    requests = [QuestionGenerationRequest(**r) for r in requests_dict]
    _run_generate_all_questions(job_id, subject_id, requests)




# Blueprint Persistence Safety Layers
from services.blueprint_repository import BlueprintRepository
from services.blueprint_loader import BlueprintLoader
from services.blueprint_guard import BlueprintGuard

# Create upload directories
BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
UPLOAD_DIR = DATA_DIR / "uploads"
SYLLABUS_DIR = UPLOAD_DIR / "syllabus"
BOOK_DIR = UPLOAD_DIR / "books"
COURSE_OUTCOMES_DIR = UPLOAD_DIR / "course_outcomes"
PAPERS_DIR = UPLOAD_DIR / "papers"
BLUEPRINTS_DIR = UPLOAD_DIR / "blueprints"
TEMP_BLUEPRINTS_DIR = BLUEPRINTS_DIR / "temp"
STUDENT_UPLOADS_DIR = UPLOAD_DIR / "student_submissions"

for directory in [DATA_DIR, UPLOAD_DIR, SYLLABUS_DIR, BOOK_DIR, COURSE_OUTCOMES_DIR, PAPERS_DIR, BLUEPRINTS_DIR, TEMP_BLUEPRINTS_DIR, STUDENT_UPLOADS_DIR]:
    directory.mkdir(parents=True, exist_ok=True)

app = FastAPI(
    title="Quest Generator API",
    version="1.0.0",
    description=(
        "The docs are public; **calling** the endpoints requires an HTTP Bearer token "
        "except login. Use the **Authorize** button below and paste the JWT returned by "
        "`POST /api/auth/login` (without the word 'Bearer')."
    ),
    openapi_tags=[
        {"name": "Auth", "description": "Login, current user, password change, user creation."},
        {"name": "Admin", "description": "User listing and approve/reject actions (admin/HOD)."},
        {"name": "Subjects", "description": "Subject CRUD, syllabus, units, topics, course outcomes, question generation."},
        {"name": "Question Banks", "description": "Question bank management per subject."},
        {"name": "Questions", "description": "Question CRUD, bulk import, parsing and images."},
        {"name": "Question Images", "description": "Upload and retrieval of question images."},
        {"name": "Question Papers", "description": "Generate, save, list, download and delete question papers."},
        {"name": "Blueprints", "description": "Exam blueprint templates (mutations are admin-only)."},
        {"name": "Search", "description": "Scoped search across subjects, questions and papers."},
        {"name": "Jobs", "description": "Background generation job status (owner/admin only)."},
        {"name": "Dashboard", "description": "Stats and recent activity, scoped to the caller."},
        {"name": "Answer Scripts", "description": "Generate/fetch/update answer scripts per paper."},
        {"name": "Evaluations", "description": "Evaluate scripts and fetch results/reports."},
        {"name": "System", "description": "Service root/health."},
    ],
)

# CORS middleware
app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_origin_regex=r"^https?://.*",
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Initialize database on startup
@app.on_event("startup")
def startup_event():
    ensure_default_blueprint()
    print("--------------------------------------------------")
    print("SYSTEM: Auth Routes are fully loaded and active.")
    print("--------------------------------------------------")

# ==================== AUTH & ADMIN ENDPOINTS ====================

def _hash_password(password: str) -> str:
    """SHA-256 hash of password"""
    return hashlib.sha256(password.encode()).hexdigest()

def _make_token(user_id: int, email: str, role: str) -> str:
    """Create a secure JWT token encoding user info"""
    from datetime import datetime, timedelta
    payload = {
        "id": user_id, 
        "email": email, 
        "role": role,
        "exp": datetime.utcnow() + timedelta(days=7)
    }
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)

def _decode_token(token: str) -> dict:
    """Decode JWT token and return payload dict"""
    try:
        return jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
    except Exception:
        return {}


# ==================== ROLE BASED ACCESS CONTROL (RBAC) ====================

VALID_ROLES = ("admin", "hod", "staff")
ROLE_ALIASES = {"advisor": "staff", "teacher": "staff"}
PUBLIC_API_PATHS = {"/api/auth/login", "/api/auth/me", "/"}  # "/" = health banner only


def _normalize_role(role: str) -> str:
    """Map legacy/unknown roles onto the three supported roles."""
    r = (role or "").strip().lower()
    r = ROLE_ALIASES.get(r, r)
    return r if r in VALID_ROLES else "staff"


def _parse_courses(raw) -> list:
    try:
        courses = json.loads(raw) if raw else []
    except (json.JSONDecodeError, TypeError):
        courses = []
    return courses if isinstance(courses, list) else []


def _fetch_auth_user(user_id: int) -> Optional[dict]:
    """Load the full user record used for role checks."""
    connection = get_db_connection()
    if not connection:
        return None
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        cursor.execute(
            f"SELECT id, email, name, role, department, courses, status FROM users WHERE id = {placeholder}",
            (user_id,)
        )
        row = cursor.fetchone()
        cursor.close()
        connection.close()
        if not row:
            return None
        user = dict(row)
        user["role"] = _normalize_role(user.get("role"))
        user["department"] = (user.get("department") or "").strip()
        user["courses"] = _parse_courses(user.get("courses"))
        return user
    except Exception:
        try:
            if connection:
                connection.close()
        except Exception:
            pass
        return None


@app.middleware("http")
async def rbac_auth_middleware(request: Request, call_next):
    """Require a valid JWT on every /api/* route (except public auth routes).
    The API docs (/docs, /redoc, /openapi.json) stay public."""
    path = request.url.path
    if request.method == "OPTIONS" or not path.startswith("/api/") or path in PUBLIC_API_PATHS:
        return await call_next(request)

    token = None
    auth_header = request.headers.get("Authorization") or ""
    if auth_header.lower().startswith("bearer "):
        token = auth_header[7:].strip()
    if not token:
        token = request.query_params.get("token")

    payload = _decode_token(token) if token else {}
    if not payload or "id" not in payload:
        return JSONResponse(status_code=401, content={"detail": "Not authenticated"}, headers={"WWW-Authenticate": "Bearer"})

    user = _fetch_auth_user(payload["id"])
    if not user:
        return JSONResponse(status_code=401, content={"detail": "Not authenticated"}, headers={"WWW-Authenticate": "Bearer"})

    request.state.auth_user = user
    return await call_next(request)


# Global security scheme: shows the "Authorize" button in /docs and lets
# Swagger UI send HTTP Bearer tokens. Enforced again by the middleware above;
# the dependency also validates when the middleware did not (public auth routes).
bearer_scheme = HTTPBearer(
    auto_error=False,
    description="JWT returned by `POST /api/auth/login`. Paste it here (without the word 'Bearer').",
)


def require_bearer(
    request: Request,
    credentials: Optional[HTTPAuthorizationCredentials] = Depends(bearer_scheme),
) -> Optional[dict]:
    existing = getattr(request.state, "auth_user", None)
    if existing:
        return existing
    if request.url.path in PUBLIC_API_PATHS:
        return None
    if not credentials or (credentials.scheme or "").lower() != "bearer":
        raise HTTPException(status_code=401, detail="Not authenticated", headers={"WWW-Authenticate": "Bearer"})
    payload = _decode_token(credentials.credentials)
    if not payload or "id" not in payload:
        raise HTTPException(status_code=401, detail="Invalid or expired token", headers={"WWW-Authenticate": "Bearer"})
    user = _fetch_auth_user(payload["id"])
    if not user:
        raise HTTPException(status_code=401, detail="Invalid or expired token", headers={"WWW-Authenticate": "Bearer"})
    request.state.auth_user = user
    return user


# Must be registered before the route definitions below so every operation
# carries the HTTPBearer security requirement in the OpenAPI schema.
app.router.dependencies.append(Depends(require_bearer))


def _current_user(request: Request) -> dict:
    user = getattr(request.state, "auth_user", None)
    if not user:
        raise HTTPException(status_code=401, detail="Not authenticated")
    return user


def _require_roles(request: Request, *roles: str) -> dict:
    """Allow only the given roles, otherwise 403."""
    user = _current_user(request)
    if user["role"] not in roles:
        raise HTTPException(status_code=403, detail="You do not have permission to perform this action")
    return user


def _assigned_subject_keys(user: dict) -> set:
    """Subject names/codes assigned to a staff member via users.courses."""
    keys = set()
    for course in user.get("courses") or []:
        if isinstance(course, dict):
            for field in ("subject", "subject_id", "name"):
                value = course.get(field)
                if value:
                    keys.add(str(value).strip().lower())
        elif isinstance(course, str) and course.strip():
            keys.add(course.strip().lower())
    return keys


def _subject_access_ids(connection, user_id: int) -> set:
    cursor = get_cursor(connection)
    placeholder = get_placeholder()
    cursor.execute(
        f"SELECT subject_id FROM subject_access WHERE user_id = {placeholder}",
        (user_id,)
    )
    rows = cursor.fetchall()
    cursor.close()
    return {row["subject_id"] for row in rows}


def _subject_visible(user: dict, subject_row: dict, connection=None, access_ids: Optional[set] = None) -> bool:
    """Check whether a subject row is within the user's scope."""
    role = user["role"]
    if role == "admin":
        return True
    if role == "hod":
        subject_dept = (subject_row.get("department") or "").strip().lower()
        user_dept = (user.get("department") or "").strip().lower()
        return bool(subject_dept) and bool(user_dept) and subject_dept == user_dept
    # staff: explicit assignments only
    if access_ids is not None:
        return subject_row.get("id") in access_ids
    if not connection:
        return False
    # fall back to checking the access map directly with connection
    cursor = get_cursor(connection)
    placeholder = get_placeholder()
    cursor.execute(
        f"SELECT 1 FROM subject_access WHERE subject_id = {placeholder} AND user_id = {placeholder} LIMIT 1",
        (subject_row.get("id"), user["id"])
    )
    row = cursor.fetchone()
    cursor.close()
    return bool(row)


def _assert_subject_access(user: dict, subject_row: dict, connection=None, access_ids: Optional[set] = None) -> None:
    if not _subject_visible(user, subject_row, connection=connection, access_ids=access_ids):
        raise HTTPException(status_code=403, detail="You do not have access to this subject")


def _visible_subject_rows(connection, user: dict) -> list:
    """All subject rows within the caller's scope (None-safe for admin)."""
    cursor = get_cursor(connection)
    cursor.execute("SELECT id, subject_id, name, department FROM subjects")
    rows = [dict(r) for r in cursor.fetchall()]
    cursor.close()
    access_ids = _subject_access_ids(connection, user["id"]) if user["role"] == "staff" else None
    return rows if user["role"] == "admin" else [r for r in rows if _subject_visible(user, r, connection=connection, access_ids=access_ids)]


def _allowed_subject_ids(connection, user: dict) -> Optional[set]:
    """None means unrestricted (admin). Otherwise a set of subject ids."""
    if user["role"] == "admin":
        return None
    return {r["id"] for r in _visible_subject_rows(connection, user)}


def _fetch_subject_row(connection, subject_id: int) -> dict:
    cursor = get_cursor(connection)
    placeholder = get_placeholder()
    cursor.execute(
        f"SELECT id, subject_id, name, department FROM subjects WHERE id = {placeholder}",
        (subject_id,)
    )
    row = cursor.fetchone()
    cursor.close()
    if not row:
        raise HTTPException(status_code=404, detail="Subject not found")
    return dict(row)


def _assert_subject_id_access(connection, user: dict, subject_id: int) -> dict:
    subject_row = _fetch_subject_row(connection, subject_id)
    _assert_subject_access(user, subject_row, connection=connection)
    return subject_row


def _assert_paper_access(connection, user: dict, paper_id: int) -> None:
    """Check access to a question paper through its subject."""
    if user["role"] == "admin":
        return
    cursor = get_cursor(connection)
    placeholder = get_placeholder()
    cursor.execute(f"SELECT subject_id FROM question_papers WHERE id = {placeholder}", (paper_id,))
    row = cursor.fetchone()
    cursor.close()
    if not row:
        raise HTTPException(status_code=404, detail="Question paper not found")
    subject_row = _fetch_subject_row(connection, row["subject_id"])
    _assert_subject_access(user, subject_row, connection=connection)


def _assert_job_access(request: Request, job_id: str) -> dict:
    user = _current_user(request)
    job = GENERATION_JOBS.get(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Job not found")
    owner_id = job.get("user_id")
    if user.get("role") != "admin" and owner_id is not None and owner_id != user["id"]:
        raise HTTPException(status_code=403, detail="Not authorized to access this job")
    return job


def _set_job_state(job_id: str, data: dict) -> None:
    existing = GENERATION_JOBS.get(job_id) or {}
    for key in ("user_id", "subject_id"):
        if key in existing and key not in data:
            data[key] = existing[key]
    GENERATION_JOBS[job_id] = data


class LoginRequest(BaseModel):
    email: str
    password: str

class CreateUserRequest(BaseModel):
    name: str
    email: str
    department: Optional[str] = None
    role: str = "staff"  # "admin", "hod" or "staff"
    password: Optional[str] = "12345678"
    courses: Optional[List[dict]] = None

class ChangePasswordRequest(BaseModel):
    token: str
    new_password: str

class AdminAction(BaseModel):
    user_id: int
    action: str  # "approve" or "reject"


class UpdateUserRole(BaseModel):
    user_id: int
    role: str


@app.post("/api/auth/login", tags=["Auth"])
def login(request: LoginRequest):
    """Authenticate user against quest_generator.db and return token + user info"""
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        cursor.execute(
            f"SELECT id, email, name, role, department, status, must_change_password, courses FROM users WHERE email = {placeholder}",
            (request.email.strip().lower(),)
        )
        user = cursor.fetchone()
        if not user:
            raise HTTPException(status_code=401, detail="Invalid email or password")

        # Verify password
        pw_hash = _hash_password(request.password.strip())
        cursor.execute(
            f"SELECT id FROM users WHERE email = {placeholder} AND password_hash = {placeholder}",
            (request.email.strip().lower(), pw_hash)
        )
        match = cursor.fetchone()
        if not match:
            raise HTTPException(status_code=401, detail="Invalid email or password")

        # Update last login
        cursor.execute(
            f"UPDATE users SET last_login = {placeholder} WHERE id = {placeholder}",
            (datetime.now(), user["id"])
        )
        connection.commit()
        cursor.close()
        connection.close()

        token = _make_token(user["id"], user["email"], _normalize_role(user["role"]))
        
        # Parse courses if available
        try:
            courses = json.loads(user["courses"]) if user["courses"] else []
        except (json.JSONDecodeError, TypeError):
            courses = []
        
        return {
            "success": True,
            "token": token,
            "user": {
                "id": user["id"],
                "email": user["email"],
                "name": user["name"],
                "role": _normalize_role(user["role"]),
                "department": user["department"],
                "must_change_password": bool(user["must_change_password"]),
                "mustChangePassword": bool(user["must_change_password"]),
                "courses": courses
            }
        }
    except HTTPException:
        raise
    except Exception as e:
        if connection: connection.close()
        raise HTTPException(status_code=500, detail=f"Login error: {str(e)}")


@app.post("/api/auth/create-user", tags=["Auth"])
def create_user(request: Request, create_request: CreateUserRequest):
    """Admin creates any user; HOD can only add staff inside their own department"""
    actor = _require_roles(request, "admin", "hod")
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        email = create_request.email.strip().lower()

        # Check duplicate
        cursor.execute(f"SELECT id FROM users WHERE email = {placeholder}", (email,))
        if cursor.fetchone():
            raise HTTPException(status_code=400, detail="User with this email already exists")

        role = create_request.role if create_request.role in VALID_ROLES else "staff"
        department = create_request.department

        if actor["role"] == "hod":
            # HODs may only add staff to their own department
            if role != "staff":
                raise HTTPException(status_code=403, detail="HOD can only add staff users")
            role = "staff"
            department = actor["department"]

        pw_hash = _hash_password(create_request.password or "12345678")
        must_change = 1 if (create_request.password is None or create_request.password == "12345678") else 0
        courses_str = json.dumps(create_request.courses) if create_request.courses else "[]"

        cursor.execute(
            f"INSERT INTO users (email, name, role, department, password_hash, status, must_change_password, courses) "
            f"VALUES ({placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, 'approved', {placeholder}, {placeholder})",
            (email, create_request.name, role, department, pw_hash, must_change, courses_str)
        )
        connection.commit()
        new_id = cursor.lastrowid
        cursor.close()
        connection.close()

        return {
            "success": True,
            "message": f"User {create_request.name} created with role '{role}'. Default password: 12345678",
            "user_id": new_id
        }
    except HTTPException:
        raise
    except Exception as e:
        if connection: connection.close()
        raise HTTPException(status_code=500, detail=f"Error creating user: {str(e)}")


@app.post("/api/auth/change-password", tags=["Auth"])
def change_password(request: ChangePasswordRequest, http_request: Request):
    """Allow a logged-in user to change their own password"""
    actor = _current_user(http_request)
    payload = _decode_token(request.token)
    if not payload or "id" not in payload:
        raise HTTPException(status_code=401, detail="Invalid token")
    if payload["id"] != actor["id"]:
        raise HTTPException(status_code=403, detail="You may only change your own password")

    if len(request.new_password) < 8:
        raise HTTPException(status_code=400, detail="Password must be at least 8 characters")

    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        pw_hash = _hash_password(request.new_password)
        cursor.execute(
            f"UPDATE users SET password_hash = {placeholder}, must_change_password = 0 WHERE id = {placeholder}",
            (pw_hash, actor["id"])
        )
        connection.commit()
        cursor.close()
        connection.close()
        return {"success": True, "message": "Password updated successfully"}
    except Exception as e:
        if connection: connection.close()
        raise HTTPException(status_code=500, detail=f"Error changing password: {str(e)}")


@app.get("/api/auth/me", tags=["Auth"])
def get_me(token: Optional[str] = None):
    """Return user info from token"""
    if not token:
        raise HTTPException(status_code=401, detail="Not authenticated")
    payload = _decode_token(token)
    if not payload or "id" not in payload:
        raise HTTPException(status_code=401, detail="Invalid token")
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        cursor.execute(
            f"SELECT id, email, name, role, department, must_change_password, courses FROM users WHERE id = {placeholder}",
            (payload["id"],)
        )
        user = cursor.fetchone()
        cursor.close()
        connection.close()
        if not user:
            raise HTTPException(status_code=404, detail="User not found")
        
        # Parse courses if available
        try:
            courses = json.loads(user["courses"]) if user["courses"] else []
        except (json.JSONDecodeError, TypeError):
            courses = []
        
        return {
            "id": user["id"],
            "email": user["email"],
            "name": user["name"],
            "role": _normalize_role(user["role"]),
            "department": user["department"],
            "must_change_password": bool(user["must_change_password"]),
            "mustChangePassword": bool(user["must_change_password"]),
            "courses": courses
        }
    except HTTPException:
        raise
    except Exception as e:
        if connection: connection.close()
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/admin/users", tags=["Admin"])
def get_all_users(request: Request):
    """Admin sees everyone; HOD only sees users of their own department"""
    actor = _require_roles(request, "admin", "hod")
    connection = get_db_connection()
    try:
        cursor = get_cursor(connection)
        if actor["role"] == "hod":
            placeholder = get_placeholder()
            cursor.execute(
                f"SELECT id, email, name, role, department, status, created_at, last_login, courses FROM users "
                f"WHERE department = {placeholder} ORDER BY created_at DESC",
                (actor["department"],)
            )
        else:
            cursor.execute(
                "SELECT id, email, name, role, department, status, created_at, last_login, courses FROM users ORDER BY created_at DESC"
            )
        users = cursor.fetchall()
        user_list = []
        for u in users:
            ud = dict(u)
            ud["role"] = _normalize_role(ud.get("role"))
            ud["courses"] = json.loads(ud["courses"]) if ud.get("courses") else []
            if ud["id"] == actor["id"]:
                continue
            if actor["role"] == "hod" and ud["role"] == "admin":
                continue
            user_list.append(ud)
        cursor.close()
        connection.close()
        return user_list
    except HTTPException:
        raise
    except Exception as e:
        if connection: connection.close()
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/admin/action", tags=["Admin"])
def admin_action(action: AdminAction, request: Request):
    """Approve or Reject a user (HOD: only inside their own department)"""
    actor = _require_roles(request, "admin", "hod")
    connection = get_db_connection()
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()

        cursor.execute(
            f"SELECT id, department FROM users WHERE id = {placeholder}",
            (action.user_id,)
        )
        target = cursor.fetchone()
        if not target:
            cursor.close()
            connection.close()
            raise HTTPException(status_code=404, detail="User not found")
        if actor["role"] == "hod" and (target.get("department") or "").strip().lower() != actor["department"].strip().lower():
            cursor.close()
            connection.close()
            raise HTTPException(status_code=403, detail="You can only manage users from your department")

        new_status = "approved" if action.action == "approve" else "rejected"
        cursor.execute(
            f"UPDATE users SET status = {placeholder} WHERE id = {placeholder}",
            (new_status, action.user_id)
        )
        connection.commit()
        cursor.close()
        connection.close()
        return {"success": True, "new_status": new_status}
    except HTTPException:
        raise
    except Exception as e:
        if connection: connection.close()
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/admin/update-role", tags=["Admin"])
def update_user_role(payload: UpdateUserRole, request: Request):
    """Allow admins to change a user's role between staff and hod."""
    actor = _require_roles(request, "admin")
    new_role = _normalize_role(payload.role)
    if new_role not in {"staff", "hod"}:
        raise HTTPException(status_code=400, detail="Role must be staff or hod")

    connection = get_db_connection()
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()

        cursor.execute(
            f"SELECT id, role, department FROM users WHERE id = {placeholder}",
            (payload.user_id,)
        )
        target = cursor.fetchone()
        if not target:
            cursor.close()
            connection.close()
            raise HTTPException(status_code=404, detail="User not found")

        if target["id"] == actor["id"]:
            cursor.close()
            connection.close()
            raise HTTPException(status_code=400, detail="You cannot change your own role")

        current_role = _normalize_role(target.get("role"))
        if current_role == new_role:
            cursor.close()
            connection.close()
            return {"success": True, "role": new_role, "message": "Role already set"}

        cursor.execute(
            f"UPDATE users SET role = {placeholder} WHERE id = {placeholder}",
            (new_role, payload.user_id)
        )
        connection.commit()
        cursor.close()
        connection.close()
        return {"success": True, "role": new_role}
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=str(e))



@app.get("/", tags=["System"])
def root():
    return {"message": "Quest Generator API", "version": "1.0.0"}

# ==================== SUBJECT ENDPOINTS ====================

@app.post("/api/subjects", tags=["Subjects"], response_model=SubjectResponse, status_code=status.HTTP_201_CREATED)
def create_subject(
    request: Request,
    subject_id: str = Form(...),
    name: str = Form(...),
    syllabus_file: Optional[UploadFile] = File(None),
    book_file: Optional[UploadFile] = File(None),
    course_outcome_file: Optional[UploadFile] = File(None),
    use_book_for_generation: bool = Form(False),
    department: Optional[str] = Form(None),
    assigned_user_ids: Optional[str] = Form(None)
):
    """Create a new subject with file uploads (admin or HOD)."""
    actor = _require_roles(request, "admin", "hod")
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        cursor.execute(f"SELECT id FROM subjects WHERE subject_id = {placeholder} OR name = {placeholder}", (subject_id, name))
        existing = cursor.fetchone()
        if existing:
            raise HTTPException(status_code=400, detail="Subject with this ID or name already exists")

        resolved_department = (department or "").strip() or None
        if actor["role"] == "hod":
            resolved_department = (actor.get("department") or "").strip() or None
        if not resolved_department and actor["role"] == "hod":
            raise HTTPException(status_code=400, detail="HOD subjects must belong to a department")

        selected_user_ids = []
        if assigned_user_ids:
            try:
                raw_ids = json.loads(assigned_user_ids)
                if isinstance(raw_ids, list):
                    selected_user_ids = [int(uid) for uid in raw_ids if str(uid).strip().isdigit()]
            except Exception:
                selected_user_ids = []
        
        safe_subject_id = "".join(c if c.isalnum() or c in "-_" else "_" for c in subject_id)
        
        syllabus_path = None
        if syllabus_file and syllabus_file.filename:
            file_extension = os.path.splitext(syllabus_file.filename)[1]
            syllabus_filename = f"{safe_subject_id}{file_extension}"
            syllabus_path = SYLLABUS_DIR / syllabus_filename
            
            with open(syllabus_path, "wb") as buffer:
                shutil.copyfileobj(syllabus_file.file, buffer)
            
            syllabus_path = str(syllabus_path)
        
        book_path = None
        if book_file and book_file.filename:
            file_extension = os.path.splitext(book_file.filename)[1]
            book_filename = f"{safe_subject_id}{file_extension}"
            book_path = BOOK_DIR / book_filename
            
            with open(book_path, "wb") as buffer:
                shutil.copyfileobj(book_file.file, buffer)
            
            book_path = str(book_path)

            if book_path.lower().endswith('.pdf'):
                try:
                    from services.image_extractor import ingest_pdf_images_to_database

                    saved_count = ingest_pdf_images_to_database(
                        book_path,
                        source_reference_prefix=f"subject:{subject_id}",
                    )
                    print(f"Book image ingestion completed. Saved images: {saved_count}")
                except Exception as img_exc:
                    print(f"Warning: Failed to ingest book images: {img_exc}")

        course_outcome_path = None
        if course_outcome_file and course_outcome_file.filename:
            file_extension = os.path.splitext(course_outcome_file.filename)[1]
            co_filename = f"{safe_subject_id}_co{file_extension}"
            course_outcome_path = COURSE_OUTCOMES_DIR / co_filename

            with open(course_outcome_path, "wb") as buffer:
                shutil.copyfileobj(course_outcome_file.file, buffer)

            course_outcome_path = str(course_outcome_path)
        
        query = f"""
            INSERT INTO subjects (subject_id, name, syllabus_file, book_file, course_outcome_file, use_book_for_generation, department)
            VALUES ({placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder})
        """
        cursor.execute(query, (subject_id, name, syllabus_path, book_path, course_outcome_path, use_book_for_generation, resolved_department))
        connection.commit()
        
        subject_db_id = cursor.lastrowid

        if selected_user_ids:
            placeholder_list = ", ".join([get_placeholder()] * len(selected_user_ids))
            cursor.execute(
                f"SELECT id, department, role FROM users WHERE id IN ({placeholder_list})",
                tuple(selected_user_ids)
            )
            eligible_users = []
            for user_row in cursor.fetchall():
                user_department = (user_row.get("department") or "").strip().lower()
                subject_department = (resolved_department or "").strip().lower()
                if actor["role"] == "hod":
                    if not subject_department or user_department != subject_department:
                        raise HTTPException(status_code=403, detail="HOD can only map users from their department")
                if user_row.get("role") == "admin":
                    continue
                eligible_users.append(user_row["id"])

            for user_id in eligible_users:
                cursor.execute(
                    f"INSERT IGNORE INTO subject_access (subject_id, user_id, granted_by) VALUES ({placeholder}, {placeholder}, {placeholder})",
                    (subject_db_id, user_id, actor["id"])
                )
            connection.commit()
        
        if syllabus_path and syllabus_path.lower().endswith('.pdf'):
            try:
                parsed_data = parse_syllabus(syllabus_path)
                if parsed_data['success']:
                    save_syllabus_to_db(connection, subject_db_id, parsed_data, subject_id)
                else:
                    print(f"Warning: Failed to parse syllabus: {parsed_data.get('error')}")
            except Exception as e:
                print(f"Warning: Error parsing syllabus: {str(e)}")
        
        cursor.execute(f"SELECT * FROM subjects WHERE id = {placeholder}", (subject_db_id,))
        result = cursor.fetchone()
        
        cursor.close()
        connection.close()
        
        return dict(result)
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.rollback()
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error creating subject: {str(e)}")

@app.get("/api/subjects", tags=["Subjects"], response_model=List[SubjectResponse])
def get_subjects(request: Request):
    """Get subjects visible to the caller (admin: all, HOD: own department, staff: assigned)"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        cursor.execute("SELECT * FROM subjects ORDER BY created_at DESC")
        subjects = cursor.fetchall()
        
        cursor.close()
        connection.close()
        
        visible = [dict(s) for s in subjects if _subject_visible(user, dict(s), connection=connection)]
        return visible
    except Exception as e:
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error fetching subjects: {str(e)}")

@app.get("/api/subjects/{subject_id}", tags=["Subjects"], response_model=SubjectResponse)
def get_subject(subject_id: int, request: Request):
    """Get a specific subject by ID"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        cursor.execute(f"SELECT * FROM subjects WHERE id = {placeholder}", (subject_id,))
        subject = cursor.fetchone()
        
        cursor.close()
        connection.close()
        
        if not subject:
            raise HTTPException(status_code=404, detail="Subject not found")
        
        _assert_subject_access(user, dict(subject), connection=connection)
        return dict(subject)
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error fetching subject: {str(e)}")


@app.post("/api/subjects/{subject_id}/course-outcome", tags=["Subjects"], response_model=SubjectResponse)
def upload_subject_course_outcome(
    request: Request,
    subject_id: int,
    course_outcome_file: UploadFile = File(...)
):
    """Upload or replace course outcome file for an existing subject"""
    actor = _require_roles(request, "admin", "hod")
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")

    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()

        cursor.execute(f"SELECT id, subject_id, department FROM subjects WHERE id = {placeholder}", (subject_id,))
        subject = cursor.fetchone()
        if not subject:
            raise HTTPException(status_code=404, detail="Subject not found")
        _assert_subject_access(actor, dict(subject), connection=connection)

        if not course_outcome_file or not course_outcome_file.filename:
            raise HTTPException(status_code=400, detail="Course outcome file is required")

        file_extension = os.path.splitext(course_outcome_file.filename)[1].lower()
        allowed_ext = {'.png', '.jpg', '.jpeg', '.pdf', '.doc', '.docx'}
        if file_extension not in allowed_ext:
            raise HTTPException(status_code=400, detail="Only image, PDF or Word files are allowed")

        raw_subject_code = subject['subject_id']
        safe_subject_id = "".join(c if c.isalnum() or c in "-_" else "_" for c in raw_subject_code)
        co_filename = f"{safe_subject_id}_co{file_extension}"
        co_path = COURSE_OUTCOMES_DIR / co_filename

        with open(co_path, "wb") as buffer:
            shutil.copyfileobj(course_outcome_file.file, buffer)

        cursor.execute(
            f"UPDATE subjects SET course_outcome_file = {placeholder} WHERE id = {placeholder}",
            (str(co_path), subject_id),
        )
        connection.commit()

        cursor.execute(f"SELECT * FROM subjects WHERE id = {placeholder}", (subject_id,))
        updated = cursor.fetchone()

        cursor.close()
        connection.close()

        return dict(updated)
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.rollback()
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error uploading course outcome: {str(e)}")


@app.get("/api/subjects/{subject_id}/course-outcome-file", tags=["Subjects"])
def download_subject_course_outcome(subject_id: int, request: Request):
    """Download the course outcome file for a subject"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")

    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        cursor.execute(
            f"SELECT course_outcome_file, subject_id, name, department FROM subjects WHERE id = {placeholder}",
            (subject_id,),
        )
        subject = cursor.fetchone()

        cursor.close()
        connection.close()

        if not subject:
            raise HTTPException(status_code=404, detail="Subject not found")

        _assert_subject_access(user, dict(subject), connection=connection)

        course_outcome_file = subject.get("course_outcome_file")
        if not course_outcome_file or not os.path.exists(course_outcome_file):
            raise HTTPException(status_code=404, detail="Course outcome file not found")

        ext = os.path.splitext(course_outcome_file)[1].lower()
        if ext in [".png", ".jpg", ".jpeg"]:
            media_type = "image/png" if ext == ".png" else "image/jpeg"
        elif ext == ".pdf":
            media_type = "application/pdf"
        elif ext in [".doc", ".docx"]:
            media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        else:
            media_type = "application/octet-stream"

        return FileResponse(
            path=course_outcome_file,
            media_type=media_type,
            filename=os.path.basename(course_outcome_file),
        )
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error downloading course outcome: {str(e)}")

@app.put("/api/subjects/{subject_id}", tags=["Subjects"], response_model=SubjectResponse)
def update_subject(subject_id: int, subject: SubjectUpdate, request: Request):
    """Update a subject (admin, or HOD of that subject's department)"""
    actor = _require_roles(request, "admin", "hod")
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        cursor.execute(f"SELECT id, subject_id, name, department FROM subjects WHERE id = {placeholder}", (subject_id,))
        existing = cursor.fetchone()
        if not existing:
            raise HTTPException(status_code=404, detail="Subject not found")
        _assert_subject_access(actor, dict(existing), connection=connection)
        
        update_fields = []
        values = []
        
        if subject.name is not None:
            update_fields.append(f"name = {placeholder}")
            values.append(subject.name)
        if subject.syllabus_file is not None:
            update_fields.append(f"syllabus_file = {placeholder}")
            values.append(subject.syllabus_file)
        if subject.book_file is not None:
            update_fields.append(f"book_file = {placeholder}")
            values.append(subject.book_file)
        if subject.course_outcome_file is not None:
            update_fields.append(f"course_outcome_file = {placeholder}")
            values.append(subject.course_outcome_file)
        if subject.use_book_for_generation is not None:
            update_fields.append(f"use_book_for_generation = {placeholder}")
            values.append(subject.use_book_for_generation)
        if getattr(subject, "department", None) is not None:
            update_fields.append(f"department = {placeholder}")
            values.append(subject.department.strip() or None)
        
        if not update_fields:
            raise HTTPException(status_code=400, detail="No fields to update")
        
        values.append(subject_id)
        query = f"UPDATE subjects SET {', '.join(update_fields)} WHERE id = {placeholder}"
        cursor.execute(query, tuple(values))
        connection.commit()
        
        cursor.execute(f"SELECT * FROM subjects WHERE id = {placeholder}", (subject_id,))
        result = cursor.fetchone()
        
        cursor.close()
        connection.close()
        
        return dict(result)
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.rollback()
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error updating subject: {str(e)}")

@app.get("/api/subjects/{subject_id}/syllabus", tags=["Subjects"])
def get_subject_syllabus(subject_id: int, http_request: Request):
    """Get parsed syllabus structure for a subject"""
    user = _current_user(http_request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        cursor.execute(f"SELECT * FROM subjects WHERE id = {placeholder}", (subject_id,))
        subject = cursor.fetchone()
        
        if not subject:
            raise HTTPException(status_code=404, detail="Subject not found")
        
        _assert_subject_access(user, dict(subject), connection=connection)
        
        cursor.execute(f"""
            SELECT * FROM units 
            WHERE subject_id = {placeholder} 
            ORDER BY unit_number
        """, (subject_id,))
        units = cursor.fetchall()
        
        result = []
        for unit in units:
            cursor.execute(f"""
                SELECT * FROM topics 
                WHERE unit_id = {placeholder}
            """, (unit['id'],))
            topics = cursor.fetchall()
            
            topics_data = []
            for topic in topics:
                cursor.execute(f"""
                    SELECT * FROM subtopics 
                    WHERE topic_id = {placeholder}
                """, (topic['id'],))
                subtopics = cursor.fetchall()
                
                topics_data.append({
                    'id': topic['id'],
                    'topic_name': topic['topic_name'],
                    'subtopics': [s['subtopic_name'] for s in subtopics]
                })
            
            result.append({
                'id': unit['id'],
                'unit_number': unit['unit_number'],
                'unit_title': unit['unit_title'],
                'topics': topics_data
            })
        
        cursor.close()
        connection.close()
        
        return {
            'subject_id': subject['subject_id'],
            'subject_name': subject['name'],
            'units': result
        }
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error fetching syllabus: {str(e)}")

@app.get("/api/subjects/{subject_id}/units", tags=["Subjects"])
def get_subject_units(subject_id: int, http_request: Request):
    """Get units for a subject"""
    user = _current_user(http_request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        cursor.execute(f"SELECT * FROM subjects WHERE id = {placeholder}", (subject_id,))
        subject = cursor.fetchone()
        
        if not subject:
            raise HTTPException(status_code=404, detail="Subject not found")
        
        _assert_subject_access(user, dict(subject), connection=connection)
        
        cursor.execute(f"""
            SELECT id, unit_number, unit_title 
            FROM units 
            WHERE subject_id = {placeholder} 
            ORDER BY unit_number
        """, (subject_id,))
        units = cursor.fetchall()
        
        cursor.close()
        connection.close()
        
        return {
            'subject_id': subject['subject_id'],
            'subject_name': subject['name'],
            'units': [dict(u) for u in units]
        }
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error fetching units: {str(e)}")

@app.get("/api/subjects/{subject_id}/topics", tags=["Subjects"])
def get_subject_topics(subject_id: int, http_request: Request, from_unit: Optional[int] = None, to_unit: Optional[int] = None):
    """Get topics and subtopics for a subject, optionally filtered by unit range.

    This endpoint now returns both top-level topics and their subtopics as
    flattened "topics" so the UI and generators can cover the full syllabus.
    """
    user = _current_user(http_request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()

        cursor.execute(f"SELECT * FROM subjects WHERE id = {placeholder}", (subject_id,))
        subject = cursor.fetchone()
        if not subject:
            raise HTTPException(status_code=404, detail="Subject not found")
        _assert_subject_access(user, dict(subject), connection=connection)
        
        # Base topics
        if from_unit is not None and to_unit is not None:
            topic_query = f"""
                SELECT t.id, t.topic_name, u.unit_number, u.unit_title
                FROM topics t
                JOIN units u ON t.unit_id = u.id
                WHERE u.subject_id = {placeholder} AND u.unit_number BETWEEN {placeholder} AND {placeholder}
                ORDER BY u.unit_number, t.id
            """
            cursor.execute(topic_query, (subject_id, from_unit, to_unit))
        else:
            topic_query = f"""
                SELECT t.id, t.topic_name, u.unit_number, u.unit_title
                FROM topics t
                JOIN units u ON t.unit_id = u.id
                WHERE u.subject_id = {placeholder}
                ORDER BY u.unit_number, t.id
            """
            cursor.execute(topic_query, (subject_id,))

        topic_rows = cursor.fetchall()

        # Related subtopics, flattened as independent topics
        if from_unit is not None and to_unit is not None:
            subtopic_query = f"""
                SELECT s.id AS subtopic_id, s.subtopic_name, u.unit_number, u.unit_title
                FROM subtopics s
                JOIN topics t ON s.topic_id = t.id
                JOIN units u ON t.unit_id = u.id
                WHERE u.subject_id = {placeholder} AND u.unit_number BETWEEN {placeholder} AND {placeholder}
                ORDER BY u.unit_number, s.id
            """
            cursor.execute(subtopic_query, (subject_id, from_unit, to_unit))
        else:
            subtopic_query = f"""
                SELECT s.id AS subtopic_id, s.subtopic_name, u.unit_number, u.unit_title
                FROM subtopics s
                JOIN topics t ON s.topic_id = t.id
                JOIN units u ON t.unit_id = u.id
                WHERE u.subject_id = {placeholder}
                ORDER BY u.unit_number, s.id
            """
            cursor.execute(subtopic_query, (subject_id,))

        subtopic_rows = cursor.fetchall()

        cursor.close()
        connection.close()

        flattened = []

        for row in topic_rows:
            d = dict(row)
            # Keep a stable string id and mark kind for debugging/extension
            d['id'] = f"topic-{d['id']}"
            d['kind'] = 'topic'
            flattened.append(d)

        for row in subtopic_rows:
            d = {
                'id': f"subtopic-{row['subtopic_id']}",
                'topic_name': row['subtopic_name'],
                'unit_number': row['unit_number'],
                'unit_title': row['unit_title'],
                'kind': 'subtopic',
            }
            flattened.append(d)

        return {'topics': flattened}
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error fetching topics: {str(e)}")

@app.get("/api/jobs/{job_id}", tags=["Jobs"])
def get_job_status(job_id: str, http_request: Request):
    job = _assert_job_access(http_request, job_id)
    return {k: v for k, v in job.items() if k != "user_id"}

@app.post("/api/jobs/{job_id}/stop", tags=["Jobs"])
@app.post("/api/jobs/{job_id}/cancel", tags=["Jobs"])
def stop_job(job_id: str, http_request: Request):
    current_job = _assert_job_access(http_request, job_id)
    if current_job.get("status") in ["completed", "failed", "cancelled", "stopped"]:
        return {"success": True, "message": f"Job is already {current_job.get('status')}", "job_id": job_id}
    
    _set_job_state(job_id, {
        "status": "cancelled",
        "error": "Process stopped by user"
    })
    return {"success": True, "message": "Job cancellation requested", "job_id": job_id}

def update_job_progress(job_id: str, completed: int, total: int, message: str = ""):
    current = GENERATION_JOBS.get(job_id, {})
    if current and current.get("status") in ("cancelled", "stopped"):
        return
    percent = int((completed / total) * 100) if total > 0 else 0
    _set_job_state(job_id, {
        "status": "pending",
        "completed": completed,
        "total": total,
        "progress": min(percent, 99),
        "message": message or f"Generated {completed} of {total} questions ({percent}%)"
    })

def run_generate_questions(job_id: str, subject_id: int, request: QuestionGenerationRequest):
    def is_cancelled():
        job = GENERATION_JOBS.get(job_id)
        return bool(job and job.get("status") in ("cancelled", "stopped"))

    try:
        if is_cancelled():
            print(f"Job {job_id} was cancelled before starting.")
            return

        connection = get_db_connection()
        if not connection:
            _set_job_state(job_id, {"status": "failed", "error": "Database connection failed"})
            return
        
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        # Get topics and subtopics between the unit range
        query = f"""
            SELECT tu.topic_name, tu.unit_number, tu.unit_title
            FROM (
                SELECT t.topic_name AS topic_name,
                       u.unit_number AS unit_number,
                       u.unit_title AS unit_title
                FROM topics t
                JOIN units u ON t.unit_id = u.id
                WHERE u.subject_id = {placeholder} AND u.unit_number BETWEEN {placeholder} AND {placeholder}
                UNION ALL
                SELECT s.subtopic_name AS topic_name,
                       u2.unit_number AS unit_number,
                       u2.unit_title AS unit_title
                FROM subtopics s
                JOIN topics t2 ON s.topic_id = t2.id
                JOIN units u2 ON t2.unit_id = u2.id
                WHERE u2.subject_id = {placeholder} AND u2.unit_number BETWEEN {placeholder} AND {placeholder}
            ) AS tu
            ORDER BY tu.unit_number
        """
        cursor.execute(query, (subject_id, request.from_unit, request.to_unit, subject_id, request.from_unit, request.to_unit))
        topics_data = cursor.fetchall()

        cursor.execute(f"SELECT subject_id FROM subjects WHERE id = {placeholder}", (subject_id,))
        subject_record = cursor.fetchone()
        
        cursor.close()
        connection.close()

        if is_cancelled():
            print(f"Job {job_id} was cancelled after database query.")
            return
        
        if request.topics and len(request.topics) > 0:
            topics = request.topics
        else:
            if not topics_data:
                _set_job_state(job_id, {"status": "failed", "error": "No topics found for the specified unit range"})
                return
            topics = [f"{t['topic_name']} (Unit {t['unit_number']})" for t in topics_data]

        if subject_record:
            raw_id = subject_record['subject_id']
            safe_subject_code = "".join(c if c.isalnum() or c in "-_" else "_" for c in raw_id)
        else:
            safe_subject_code = str(subject_id)
        
        try:
            sync_subject_files_to_qdrant(safe_subject_code)
            query_str = f"Topics: {', '.join(topics)}"
            retrieved_data = retrieve_context(query_str, subject_id=safe_subject_code)
            context_str = format_context_for_prompt(retrieved_data)
        except Exception as rag_err:
            print(f"RAG Error (continuing without RAG): {rag_err}")
            context_str = None

        if is_cancelled():
            print(f"Job {job_id} was cancelled after RAG sync.")
            return

        total_needed = sum(item.count for item in request.plan) if request.plan else request.count
        total_needed += max(0, int(request.image_questions or 0))
        update_job_progress(job_id, 0, total_needed, f"Generating {total_needed} questions...")

        all_questions = []
        if request.plan:
            for item in request.plan:
                if is_cancelled():
                    print(f"Job {job_id} was cancelled during plan loop.")
                    return
                unit_topics = [
                    f"{t['topic_name']} (Unit {t['unit_number']})"
                    for t in topics_data
                    if int(t['unit_number']) == int(item.unit)
                ]
                if not unit_topics:
                    continue
                
                base_completed = len(all_questions)
                def on_progress(batch_completed, batch_target):
                    update_job_progress(job_id, base_completed + batch_completed, total_needed)

                unit_questions = generate_questions_with_ollama(
                    topics=unit_topics,
                    count=item.count,
                    marks=request.marks,
                    difficulty=item.difficulty,
                    part_name=request.part_name,
                    context=context_str,
                    ai_provider=request.ai_provider,
                    blooms_level=item.blooms_level,
                    cancel_check=is_cancelled,
                    progress_callback=on_progress,
                )
                for q in unit_questions:
                    q.setdefault('unit', str(item.unit))
                    q.setdefault('difficulty', item.difficulty)
                all_questions.extend(unit_questions)
                update_job_progress(job_id, len(all_questions), total_needed)
        else:
            def on_progress(batch_completed, batch_target):
                update_job_progress(job_id, batch_completed, total_needed)

            all_questions = generate_questions_with_ollama(
                topics=topics,
                count=request.count,
                marks=request.marks,
                difficulty=request.difficulty,
                part_name=request.part_name,
                context=context_str,
                ai_provider=request.ai_provider,
                cancel_check=is_cancelled,
                progress_callback=on_progress,
            )

        # Generate additional questions FROM images (web-search / user-upload / book)
        image_question_count = max(0, int(request.image_questions or 0))
        if image_question_count > 0 and not is_cancelled():
            print(f"🖼️ Generating up to {image_question_count} image-based questions...")
            try:
                from services.image_question_generator import generate_image_questions, encode_image_data_for_json

                def img_progress(batch_completed, batch_target):
                    update_job_progress(job_id, len(all_questions) + batch_completed, total_needed)

                img_questions = generate_image_questions(
                    topics=topics,
                    count=image_question_count,
                    marks=request.marks,
                    difficulty=request.difficulty,
                    part_name=request.part_name,
                    image_sources=request.image_sources,
                    ai_provider=request.ai_provider,
                    context=context_str,
                    blooms_level=None,
                    cancel_check=is_cancelled,
                    progress_callback=img_progress,
                )
                all_questions.extend(encode_image_data_for_json(q) for q in img_questions)
                update_job_progress(job_id, len(all_questions), total_needed)
            except Exception as img_err:
                print(f"❌ Image-based question generation failed: {img_err}")
                import traceback
                traceback.print_exc()

        if is_cancelled():
            print(f"Job {job_id} was cancelled before finalizing.")
            return

        _set_job_state(job_id, {
            "status": "completed",
            "completed": len(all_questions),
            "total": total_needed,
            "progress": 100,
            "result": {
                'success': True,
                'count': len(all_questions),
                'questions': all_questions,
                'topics_covered': len(topics)
            }
        })
    except Exception as e:
        if not is_cancelled():
            _set_job_state(job_id, {
                "status": "failed",
                "error": str(e)
            })

@app.post("/api/subjects/{subject_id}/generate-questions", tags=["Subjects"])
def generate_questions(subject_id: int, request: QuestionGenerationRequest, background_tasks: BackgroundTasks, http_request: Request):
    """Generate questions using Ollama based on topics from database"""
    user = _current_user(http_request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    try:
        _assert_subject_id_access(connection, user, subject_id)
    finally:
        connection.close()

    if not test_ollama_connection(request.ai_provider):
        raise HTTPException(
            status_code=503, 
            detail="No AI provider is reachable. Set request.ai_provider as ollama/xai/openai/gemini or configure AI_MODE in backend/.env."
        )
    
    job_id = str(uuid.uuid4())
    GENERATION_JOBS[job_id] = {"status": "pending", "user_id": user["id"], "subject_id": subject_id}
    
    if USE_CELERY and redis_client:
        celery_run_generate_questions.delay(job_id, subject_id, request.dict())
    else:
        background_tasks.add_task(run_generate_questions, job_id, subject_id, request)
        
    return {"success": True, "job_id": job_id}


@app.post("/api/subjects/{subject_id}/generate-all-questions", tags=["Subjects"])
def generate_all_questions(subject_id: int, requests: List[QuestionGenerationRequest], background_tasks: BackgroundTasks, http_request: Request):
    user = _current_user(http_request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    try:
        _assert_subject_id_access(connection, user, subject_id)
    finally:
        connection.close()

    job_id = str(uuid.uuid4())
    GENERATION_JOBS[job_id] = {"status": "pending", "user_id": user["id"], "subject_id": subject_id}
    
    if USE_CELERY and redis_client:
        celery_run_generate_all_questions.delay(job_id, subject_id, [r.dict() for r in requests])
    else:
        background_tasks.add_task(_run_generate_all_questions, job_id, subject_id, requests)
        
    return {"success": True, "job_id": job_id}

def _run_generate_all_questions(job_id: str, subject_id: int, requests: List[QuestionGenerationRequest]):
    """Generate questions for all parts at once"""
    def is_cancelled():
        job = GENERATION_JOBS.get(job_id)
        return bool(job and job.get("status") in ("cancelled", "stopped"))

    if is_cancelled():
        print(f"Job {job_id} was cancelled before starting.")
        return

    import threading
    progress_lock = threading.Lock()
    completed_counter = [0]

    grand_total = 0
    for req in requests:
        if req.plan:
            grand_total += sum(item.count for item in req.plan)
        else:
            grand_total += req.count
        grand_total += max(0, int(req.image_questions or 0))

    def increment_progress(delta_count):
        with progress_lock:
            completed_counter[0] += delta_count
            update_job_progress(job_id, completed_counter[0], grand_total)

    update_job_progress(job_id, 0, grand_total, f"Generating {grand_total} questions across all parts...")

    default_provider = requests[0].ai_provider if requests else None
    if not test_ollama_connection(default_provider):
        raise HTTPException(
            status_code=503, 
            detail="No AI provider is reachable. Set request.ai_provider as ollama/xai/openai/gemini or configure AI_MODE in backend/.env."
        )
    
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        # Get subject code for RAG once
        cursor.execute(f"SELECT subject_id FROM subjects WHERE id = {placeholder}", (subject_id,))
        subject_record = cursor.fetchone()
        if subject_record:
            raw_id = subject_record['subject_id']
            safe_subject_code = "".join(c if c.isalnum() or c in "-_" else "_" for c in raw_id)
        else:
            safe_subject_code = str(subject_id)

        # Sync files to Qdrant once for the whole request
        try:
            sync_subject_files_to_qdrant(safe_subject_code)
        except Exception as rag_err:
            print(f"RAG Sync Error: {rag_err}")

        tasks_inputs = []

        for request in requests:
            if is_cancelled():
                print(f"Job {job_id} cancelled during request preparation.")
                return

            topics_data = []

            # Determine topics for this request
            if request.topics and len(request.topics) > 0:
                topics = request.topics
            else:
                query = f"""
                    SELECT t.topic_name, u.unit_number, u.unit_title
                    FROM topics t
                    JOIN units u ON t.unit_id = u.id
                    WHERE u.subject_id = {placeholder} AND u.unit_number BETWEEN {placeholder} AND {placeholder}
                    ORDER BY u.unit_number
                """
                cursor.execute(query, (subject_id, request.from_unit, request.to_unit))
                topics_data = cursor.fetchall()
                if topics_data:
                    topics = [f"{dict(row)['topic_name']} (Unit {dict(row)['unit_number']})" for row in topics_data]
                else:
                    topics = []
            
            if topics:
                # RAG INTEGRATION: Try to get context for each part
                context_str = None
                try:
                    query_str = f"Topics: {', '.join(topics)}"
                    retrieved_data = retrieve_context(query_str, subject_id=safe_subject_code)
                    context_str = format_context_for_prompt(retrieved_data)
                except Exception as rag_err:
                    print(f"RAG Retrieval Error: {rag_err}")

                tasks_inputs.append({
                    'topics_data': topics_data,
                    'topics': topics,
                    'request': request,
                    'context_str': context_str
                })

        def fetch_questions(input_data):
            if is_cancelled():
                return {
                    'part_name': input_data['request'].part_name,
                    'success': False,
                    'error': 'Cancelled by user'
                }
            req = input_data['request']
            prev_batch_completed = [0]

            def on_progress(batch_completed, batch_target):
                delta = batch_completed - prev_batch_completed[0]
                if delta > 0:
                    prev_batch_completed[0] = batch_completed
                    increment_progress(delta)

            try:
                all_questions = []
                topics_data = input_data['topics_data']
                topics = input_data['topics']
                context_str = input_data['context_str']

                if req.plan:
                    for item in req.plan:
                        if is_cancelled():
                            break
                        unit_topics = [
                            f"{row['topic_name']} (Unit {row['unit_number']})"
                            for row in topics_data
                            if int(row['unit_number']) == int(item.unit)
                        ]

                        if not unit_topics:
                            continue

                        plan_item_prev = [0]
                        def on_plan_progress(batch_completed, batch_target):
                            delta = batch_completed - plan_item_prev[0]
                            if delta > 0:
                                plan_item_prev[0] = batch_completed
                                increment_progress(delta)

                        unit_questions = generate_questions_with_ollama(
                            topics=unit_topics,
                            count=item.count,
                            marks=req.marks,
                            difficulty=item.difficulty,
                            part_name=req.part_name,
                            context=context_str,
                            ai_provider=req.ai_provider,
                            blooms_level=item.blooms_level,
                            cancel_check=is_cancelled,
                            progress_callback=on_plan_progress,
                        )

                        for q in unit_questions:
                            q.setdefault('unit', str(item.unit))
                            q.setdefault('difficulty', item.difficulty)
                        all_questions.extend(unit_questions)
                else:
                    all_questions = generate_questions_with_ollama(
                        topics=topics,
                        count=req.count,
                        marks=req.marks,
                        difficulty=req.difficulty,
                        part_name=req.part_name,
                        context=context_str,
                        ai_provider=req.ai_provider,
                        cancel_check=is_cancelled,
                        progress_callback=on_progress,
                    )

                # Additional questions generated FROM images for this request
                image_question_count = max(0, int(req.image_questions or 0))
                if image_question_count > 0 and not is_cancelled():
                    try:
                        from services.image_question_generator import generate_image_questions, encode_image_data_for_json

                        def img_progress(batch_completed, batch_target):
                            delta = batch_completed - prev_batch_completed[0]
                            if delta > 0:
                                prev_batch_completed[0] = batch_completed
                                increment_progress(delta)

                        img_questions = generate_image_questions(
                            topics=topics,
                            count=image_question_count,
                            marks=req.marks,
                            difficulty=req.difficulty,
                            part_name=req.part_name,
                            image_sources=req.image_sources,
                            ai_provider=req.ai_provider,
                            context=context_str,
                            blooms_level=None,
                            cancel_check=is_cancelled,
                            progress_callback=img_progress,
                        )
                        all_questions.extend(encode_image_data_for_json(q) for q in img_questions)
                    except Exception as img_err:
                        print(f"❌ Image-based question generation failed for {req.part_name}: {img_err}")
                        import traceback
                        traceback.print_exc()

                return {
                    'part_name': req.part_name,
                    'success': True,
                    'count': len(all_questions),
                    'questions': all_questions,
                    'topics_covered': len(topics)
                }
            except Exception as e:
                return {
                    'part_name': req.part_name,
                    'success': False,
                    'error': str(e)
                }

        if tasks_inputs and not is_cancelled():
            import concurrent.futures
            with concurrent.futures.ThreadPoolExecutor(max_workers=2) as executor:
                all_results = list(executor.map(fetch_questions, tasks_inputs))
        else:
            all_results = []
        
        cursor.close()
        connection.close()
        
        if is_cancelled():
            print(f"Job {job_id} was cancelled before completing.")
            return

        _set_job_state(job_id, {
            "status": "completed",
            "completed": completed_counter[0],
            "total": grand_total,
            "progress": 100,
            "result": {
                'success': True,
                'parts': all_results,
                'total_parts': len(requests)
            }
        })
    
    except HTTPException as e:
        if not is_cancelled():
            _set_job_state(job_id, {
                "status": "failed",
                "error": str(e.detail)
            })
    except Exception as e:
        if connection:
            connection.close()
        if not is_cancelled():
            _set_job_state(job_id, {
                "status": "failed",
                "error": f"Error generating questions: {str(e)}"
            })


    
@app.delete("/api/subjects/{subject_id}", tags=["Subjects"], status_code=status.HTTP_204_NO_CONTENT)
def delete_subject(subject_id: int, request: Request):
    """Delete a subject and cascade delete all related data (question banks, questions, papers)"""
    _require_roles(request, "admin")
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        # Get subject details
        cursor.execute(f"""
            SELECT id, syllabus_file, book_file 
            FROM subjects 
            WHERE id = {placeholder}
        """, (subject_id,))
        subject = cursor.fetchone()
        
        if not subject:
            cursor.close()
            connection.close()
            raise HTTPException(status_code=404, detail="Subject not found")
        
        # Get all question banks for this subject (for logging/info)
        cursor.execute(f"""
            SELECT id, name 
            FROM question_banks 
            WHERE subject_id = {placeholder}
        """, (subject_id,))
        question_banks = cursor.fetchall()
        
        # Get all question paper files for this subject
        cursor.execute(f"""
            SELECT DISTINCT file_path 
            FROM question_papers 
            WHERE subject_id = {placeholder} AND file_path IS NOT NULL
        """, (subject_id,))
        paper_files = [row['file_path'] for row in cursor.fetchall()]
        
        # Delete the subject (cascade will handle question_banks, questions, etc.)
        cursor.execute(f"DELETE FROM subjects WHERE id = {placeholder}", (subject_id,))
        connection.commit()
        
        # Collect files to delete
        files_to_delete = []
        
        if subject['syllabus_file']:
            files_to_delete.append(subject['syllabus_file'])
        
        if subject['book_file']:
            files_to_delete.append(subject['book_file'])
        
        files_to_delete.extend(paper_files)
        
        # Delete files
        for file_path in files_to_delete:
            if file_path and os.path.exists(file_path):
                try:
                    os.remove(file_path)
                except Exception as e:
                    print(f"Warning: Could not delete file {file_path}: {str(e)}")
        
        # Log deletion info
        if question_banks:
            qb_names = ', '.join([qb['name'] for qb in question_banks])
            print(f"✓ Deleted subject {subject_id} and {len(question_banks)} question bank(s): {qb_names}")
        else:
            print(f"✓ Deleted subject {subject_id}")
        
        cursor.close()
        connection.close()
        
        return None
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.rollback()
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error deleting subject: {str(e)}")

# ==================== QUESTION BANK ENDPOINTS ====================

@app.post("/api/question-banks", tags=["Question Banks"], response_model=QuestionBankResponse, status_code=status.HTTP_201_CREATED)
def create_question_bank(question_bank: QuestionBankCreate, request: Request):
    """Create a new question bank"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        cursor.execute(f"SELECT id FROM subjects WHERE id = {placeholder}", (question_bank.subject_id,))
        if not cursor.fetchone():
            raise HTTPException(status_code=404, detail="Subject not found")
        _assert_subject_id_access(connection, user, question_bank.subject_id)
        
        query = f"""
            INSERT INTO question_banks (name, subject_id, description, total_questions)
            VALUES ({placeholder}, {placeholder}, {placeholder}, 0)
        """
        cursor.execute(query, (question_bank.name, question_bank.subject_id, question_bank.description))
        connection.commit()
        
        bank_id = cursor.lastrowid
        
        cursor.execute(f"SELECT * FROM question_banks WHERE id = {placeholder}", (bank_id,))
        result = cursor.fetchone()
        
        cursor.close()
        connection.close()
        
        return dict(result)
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error creating question bank: {str(e)}")

@app.get("/api/question-banks", tags=["Question Banks"], response_model=List[QuestionBankResponse])
def get_all_question_banks(request: Request):
    """Get question banks visible to the caller"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        allowed = _allowed_subject_ids(connection, user)
        cursor = get_cursor(connection)
        cursor.execute("SELECT * FROM question_banks ORDER BY created_at DESC")
        banks = cursor.fetchall()
        
        cursor.close()
        connection.close()
        
        if allowed is None:
            return [dict(bank) for bank in banks]
        return [dict(bank) for bank in banks if bank.get("subject_id") in allowed]
    except Exception as e:
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error fetching question banks: {str(e)}")

@app.get("/api/question-banks/subject/{subject_id}", tags=["Question Banks"], response_model=List[QuestionBankResponse])
def get_question_banks_by_subject(subject_id: int, request: Request):
    """Get all question banks for a specific subject"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        _assert_subject_id_access(connection, user, subject_id)
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        cursor.execute(
            f"SELECT * FROM question_banks WHERE subject_id = {placeholder} ORDER BY created_at DESC",
            (subject_id,)
        )
        banks = cursor.fetchall()
        
        cursor.close()
        connection.close()
        
        return [dict(bank) for bank in banks]
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error fetching question banks: {str(e)}")

@app.delete("/api/question-banks/{bank_id}", tags=["Question Banks"])
def delete_question_bank(bank_id: int, request: Request):
    """Delete a question bank and all its questions"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        cursor.execute(f"SELECT id, subject_id FROM question_banks WHERE id = {placeholder}", (bank_id,))
        bank_row = cursor.fetchone()
        if not bank_row:
            cursor.close()
            connection.close()
            raise HTTPException(status_code=404, detail="Question bank not found")
        _assert_subject_id_access(connection, user, bank_row["subject_id"])
        
        cursor.execute(f"DELETE FROM question_banks WHERE id = {placeholder}", (bank_id,))
        connection.commit()
        
        cursor.close()
        connection.close()
        
        return {"success": True, "message": "Question bank deleted successfully"}
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error deleting question bank: {str(e)}")

# ==================== QUESTION ENDPOINTS ====================

@app.post("/api/questions/batch", tags=["Questions"], status_code=status.HTTP_201_CREATED)
def create_questions_batch(questions: List[QuestionCreate], request: Request):
    """Create multiple questions at once"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        if not questions:
            raise HTTPException(status_code=400, detail="No questions provided")
        
        checked_subjects = set()
        checked_banks = set()
        checked_images = set()
        for question in questions:
            if question.subject_id not in checked_subjects:
                checked_subjects.add(question.subject_id)
                cursor.execute(f"SELECT id FROM subjects WHERE id = {placeholder}", (question.subject_id,))
                if not cursor.fetchone():
                    raise HTTPException(status_code=404, detail="Subject not found")
                _assert_subject_id_access(connection, user, question.subject_id)
            if question.question_bank_id not in checked_banks:
                checked_banks.add(question.question_bank_id)
                cursor.execute(
                    f"SELECT id, subject_id FROM question_banks WHERE id = {placeholder}",
                    (question.question_bank_id,)
                )
                bank = cursor.fetchone()
                if not bank:
                    raise HTTPException(status_code=404, detail="Question bank not found")
                _assert_subject_id_access(connection, user, bank["subject_id"])
            if question.image_id and question.image_id not in checked_images:
                checked_images.add(question.image_id)
                cursor.execute(
                    f"SELECT id, subject_id FROM question_images WHERE id = {placeholder}",
                    (question.image_id,)
                )
                image_row = cursor.fetchone()
                if not image_row:
                    raise HTTPException(status_code=404, detail="Image not found")
                if image_row["subject_id"] is not None:
                    _assert_subject_id_access(connection, user, image_row["subject_id"])
        
        question_ids = []
        for question in questions:
            query = f"""
                INSERT INTO questions (question_bank_id, subject_id, content, part, unit, topic, difficulty, marks, blooms_level, source, image_id)
                VALUES ({placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder})
            """
            values = (
                question.question_bank_id,
                question.subject_id,
                question.content,
                question.part,
                question.unit,
                question.topic,
                question.difficulty,
                question.marks,
                question.blooms_level,
                question.source or "teacher",
                question.image_id
            )
            cursor.execute(query, values)
            question_ids.append(cursor.lastrowid)
        
        cursor.execute(
            f"UPDATE question_banks SET total_questions = total_questions + {placeholder} WHERE id = {placeholder}",
            (len(questions), questions[0].question_bank_id)
        )
        
        connection.commit()
        
        placeholders = ','.join([placeholder] * len(question_ids))
        cursor.execute(f"SELECT * FROM questions WHERE id IN ({placeholders})", tuple(question_ids))
        results = cursor.fetchall()
        
        cursor.close()
        connection.close()
        
        return {
            'success': True,
            'count': len(results),
            'question_bank_id': questions[0].question_bank_id,
            'questions': [dict(r) for r in results]
        }
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.rollback()
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error creating questions: {str(e)}")

@app.post("/api/questions", tags=["Questions"], response_model=QuestionResponse, status_code=status.HTTP_201_CREATED)
def create_question(question: QuestionCreate, request: Request):
    """Create a new question"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        cursor.execute(f"SELECT id FROM subjects WHERE id = {placeholder}", (question.subject_id,))
        if not cursor.fetchone():
            raise HTTPException(status_code=404, detail="Subject not found")
        _assert_subject_id_access(connection, user, question.subject_id)
        
        cursor.execute(
            f"SELECT id, subject_id FROM question_banks WHERE id = {placeholder}",
            (question.question_bank_id,)
        )
        bank = cursor.fetchone()
        if not bank:
            raise HTTPException(status_code=404, detail="Question bank not found")
        _assert_subject_id_access(connection, user, bank["subject_id"])
        
        if question.image_id:
            cursor.execute(
                f"SELECT id, subject_id FROM question_images WHERE id = {placeholder}",
                (question.image_id,)
            )
            image_row = cursor.fetchone()
            if not image_row:
                raise HTTPException(status_code=404, detail="Image not found")
            if image_row["subject_id"] is not None:
                _assert_subject_id_access(connection, user, image_row["subject_id"])
        
        query = f"""
            INSERT INTO questions (question_bank_id, subject_id, content, part, unit, topic, difficulty, marks, blooms_level, source, image_id)
            VALUES ({placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder})
        """
        values = (
            question.question_bank_id,
            question.subject_id,
            question.content,
            question.part,
            question.unit,
            question.topic,
            question.difficulty,
            question.marks,
            question.blooms_level,
            question.source or "teacher",
            question.image_id
        )
        cursor.execute(query, values)
        connection.commit()
        
        question_id = cursor.lastrowid
        cursor.execute(f"SELECT * FROM questions WHERE id = {placeholder}", (question_id,))
        result = cursor.fetchone()
        
        cursor.close()
        connection.close()
        
        return dict(result)
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.rollback()
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error creating question: {str(e)}")

@app.post("/api/questions/upload-parse", tags=["Questions"])
def upload_and_parse_questions(
    file: UploadFile = File(...)
):
    """
    Parse uploaded question file (JSON, CSV, TXT, DOCX, PDF) and extract structured questions for review.
    """
    filename = file.filename or ""
    ext = filename.split(".")[-1].lower()
    content_bytes = file.file.read()
    
    parsed_questions = []
    
    try:
        if ext == "json":
            data = json.loads(content_bytes.decode("utf-8"))
            if isinstance(data, list):
                items = data
            elif isinstance(data, dict) and "questions" in data:
                items = data["questions"]
            else:
                items = [data]
            
            for item in items:
                if isinstance(item, dict) and "content" in item:
                    parsed_questions.append({
                        "content": str(item.get("content", "")).strip(),
                        "part": str(item.get("part", "Part A")).strip(),
                        "unit": str(item.get("unit", "1")).strip(),
                        "topic": str(item.get("topic", "")).strip(),
                        "difficulty": str(item.get("difficulty", "medium")).lower(),
                        "marks": float(item.get("marks", 2.0)),
                        "blooms_level": item.get("blooms_level") or item.get("bloomsLevel") or None,
                        "source": "teacher"
                    })
        elif ext == "csv":
            import csv
            text_str = content_bytes.decode("utf-8", errors="ignore")
            reader = csv.DictReader(io.StringIO(text_str))
            for row in reader:
                content = row.get("content") or row.get("question") or row.get("Question")
                if content:
                    parsed_questions.append({
                        "content": content.strip(),
                        "part": (row.get("part") or "Part A").strip(),
                        "unit": (row.get("unit") or "1").strip(),
                        "topic": (row.get("topic") or "").strip(),
                        "difficulty": (row.get("difficulty") or "medium").strip().lower(),
                        "marks": float(row.get("marks") or 2.0),
                        "blooms_level": row.get("blooms_level") or None,
                        "source": "teacher"
                    })
        elif ext in ["txt", "docx", "doc", "pdf"]:
            text_str = ""
            if ext == "txt":
                text_str = content_bytes.decode("utf-8", errors="ignore")
            elif ext in ["docx", "doc"]:
                from docx import Document
                doc = Document(io.BytesIO(content_bytes))
                text_str = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            elif ext == "pdf":
                import fitz
                pdf_doc = fitz.open(stream=content_bytes, filetype="pdf")
                for page in pdf_doc:
                    text_str += page.get_text() + "\n"
                pdf_doc.close()
            
            lines = text_str.splitlines()
            current_q = []
            for line in lines:
                stripped = line.strip()
                if not stripped:
                    continue
                if re.match(r"^(?:Q\d+[\.\)]|\d+[\.\)])\s+", stripped) and current_q:
                    q_text = " ".join(current_q).strip()
                    q_text_clean = re.sub(r"^(?:Q\d+[\.\)]|\d+[\.\)])\s*", "", q_text)
                    if q_text_clean:
                        parsed_questions.append({
                            "content": q_text_clean,
                            "part": "Part A",
                            "unit": "1",
                            "topic": "",
                            "difficulty": "medium",
                            "marks": 2.0,
                            "blooms_level": None,
                            "source": "teacher"
                        })
                    current_q = [stripped]
                else:
                    current_q.append(stripped)
            if current_q:
                q_text = " ".join(current_q).strip()
                q_text_clean = re.sub(r"^(?:Q\d+[\.\)]|\d+[\.\)])\s*", "", q_text)
                if q_text_clean:
                    parsed_questions.append({
                        "content": q_text_clean,
                        "part": "Part A",
                        "unit": "1",
                        "topic": "",
                        "difficulty": "medium",
                        "marks": 2.0,
                        "blooms_level": None,
                        "source": "teacher"
                    })
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported file format: {ext}. Supported formats: JSON, CSV, TXT, DOCX, PDF")
            
        return {
            "success": True,
            "filename": filename,
            "count": len(parsed_questions),
            "questions": parsed_questions
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error parsing questions file: {str(e)}")

@app.post("/api/question-images/upload", tags=["Question Images"])
def upload_user_image(
    request: Request,
    file: UploadFile = File(...),
    keywords: str = Form(...),
    description: Optional[str] = Form("User uploaded image"),
    subject_id: Optional[int] = Form(None),
    unit: Optional[str] = Form(None)
):
    """
    Upload teacher/user image for question paper generation
    """
    user = _current_user(request)
    if subject_id is not None:
        connection = get_db_connection()
        if not connection:
            raise HTTPException(status_code=500, detail="Database connection failed")
        try:
            _assert_subject_id_access(connection, user, subject_id)
        finally:
            connection.close()
    try:
        content_bytes = file.file.read()
        if not content_bytes:
            raise HTTPException(status_code=400, detail="Uploaded image file is empty")
        
        from services.image_service import ImageService
        
        image_id = ImageService.save_image(
            keywords=keywords,
            description=description or "User uploaded image",
            image_blob=content_bytes,
            source_type="user_uploaded",
            source_reference=f"user_upload:{file.filename}",
            file_name=file.filename,
            subject_id=subject_id,
            unit=unit
        )
        
        if not image_id:
            raise HTTPException(status_code=500, detail="Failed to save image to database")
            
        return {
            "success": True,
            "image_id": image_id,
            "file_name": file.filename,
            "keywords": keywords,
            "description": description,
            "source_type": "user_uploaded"
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error uploading image: {str(e)}")

@app.get("/api/questions/subject/{subject_id}", tags=["Questions"], response_model=List[QuestionResponse])
def get_questions_by_subject(subject_id: int, request: Request):
    """Get all questions for a specific subject"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        _assert_subject_id_access(connection, user, subject_id)
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        cursor.execute(
            f"SELECT * FROM questions WHERE subject_id = {placeholder} ORDER BY created_at DESC",
            (subject_id,)
        )
        questions = cursor.fetchall()
        
        cursor.close()
        connection.close()
        
        return [dict(q) for q in questions]
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error fetching questions: {str(e)}")

@app.get("/api/questions/bank/{bank_id}", tags=["Questions"], response_model=List[QuestionResponse])
def get_questions_by_bank(bank_id: int, request: Request):
    """Get all questions for a specific question bank"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        cursor.execute(
            f"SELECT id, name, subject_id FROM question_banks WHERE id = {placeholder}",
            (bank_id,)
        )
        bank = cursor.fetchone()
        
        if not bank:
            cursor.close()
            connection.close()
            raise HTTPException(status_code=404, detail=f"Question bank with ID {bank_id} not found")
        
        _assert_subject_id_access(connection, user, bank["subject_id"])
        
        cursor.execute(
            f"SELECT * FROM questions WHERE question_bank_id = {placeholder} ORDER BY created_at DESC",
            (bank_id,)
        )
        questions = cursor.fetchall()
        
        cursor.close()
        connection.close()
        
        if not questions:
            raise HTTPException(
                status_code=404, 
                detail=f"No questions found in question bank '{bank['name']}'. Please add questions first."
            )
        
        return [dict(q) for q in questions]
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error fetching questions: {str(e)}")

@app.get("/api/questions/by-question-bank/{question_bank_id}", tags=["Questions"], response_model=List[QuestionResponse])
def get_questions_by_question_bank_id(question_bank_id: int, request: Request):
    """Get all questions for a specific question bank - alternative endpoint"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        cursor.execute(
            f"SELECT id, name, subject_id FROM question_banks WHERE id = {placeholder}",
            (question_bank_id,)
        )
        bank = cursor.fetchone()
        
        if not bank:
            cursor.close()
            connection.close()
            raise HTTPException(status_code=404, detail=f"Question bank with ID {question_bank_id} not found")
        
        _assert_subject_id_access(connection, user, bank["subject_id"])
        
        cursor.execute(
            f"SELECT * FROM questions WHERE question_bank_id = {placeholder} ORDER BY created_at DESC",
            (question_bank_id,)
        )
        questions = cursor.fetchall()
        
        cursor.close()
        connection.close()
        
        if not questions:
            raise HTTPException(
                status_code=404, 
                detail=f"No questions found in question bank '{bank['name']}'. Please add questions first."
            )
        
        return [dict(q) for q in questions]
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error fetching questions: {str(e)}")

@app.delete("/api/questions/{question_id}", tags=["Questions"])
def delete_question(question_id: int, request: Request):
    """Delete a question by ID"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        cursor.execute(f"SELECT id, subject_id FROM questions WHERE id = {placeholder}", (question_id,))
        question = cursor.fetchone()
        
        if not question:
            cursor.close()
            connection.close()
            raise HTTPException(status_code=404, detail="Question not found")
        _assert_subject_id_access(connection, user, question["subject_id"])
        
        cursor.execute(f"DELETE FROM questions WHERE id = {placeholder}", (question_id,))
        connection.commit()
        
        cursor.close()
        connection.close()
        
        return {"success": True, "message": "Question deleted successfully"}
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error deleting question: {str(e)}")

# ==================== SEARCH ENDPOINTS ====================

@app.get("/api/search/questions", tags=["Search"])
def search_questions(
    request: Request,
    q: str = "",
    subject_id: Optional[int] = None,
    bank_id: Optional[int] = None,
    difficulty: Optional[str] = None,
    unit: Optional[str] = None,
    limit: int = 50
):
    """Advanced search for questions with filters"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        allowed = _allowed_subject_ids(connection, user)
        if allowed is not None and not allowed:
            return {
                "success": True,
                "count": 0,
                "results": [],
                "query": q,
                "filters": {
                    "subject_id": subject_id,
                    "bank_id": bank_id,
                    "difficulty": difficulty,
                    "unit": unit
                }
            }
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        # Build dynamic query
        query = "SELECT id, content, unit, topic, difficulty, marks, part, question_bank_id, subject_id, created_at FROM questions WHERE 1=1"
        params = []
        
        # Search term - search in content, topic, and unit
        if q:
            query += f" AND (content LIKE {placeholder} OR topic LIKE {placeholder} OR unit LIKE {placeholder})"
            search_term = f"%{q}%"
            params.extend([search_term, search_term, search_term])
        
        # Subject filter
        if subject_id:
            query += f" AND subject_id = {placeholder}"
            params.append(subject_id)
        
        # Bank filter
        if bank_id:
            query += f" AND question_bank_id = {placeholder}"
            params.append(bank_id)
        
        # Difficulty filter
        if difficulty:
            query += f" AND difficulty = {placeholder}"
            params.append(difficulty)
        
        # Unit filter
        if unit:
            query += f" AND unit = {placeholder}"
            params.append(unit)
        
        if allowed is not None:
            id_placeholders = ','.join([placeholder] * len(allowed))
            query += f" AND subject_id IN ({id_placeholders})"
            params.extend(list(allowed))
        
        query += f" ORDER BY created_at DESC LIMIT {placeholder}"
        params.append(limit)
        
        cursor.execute(query, params)
        results = cursor.fetchall()
        cursor.close()
        connection.close()
        
        return {
            "success": True,
            "count": len(results),
            "results": [dict(r) for r in results],
            "query": q,
            "filters": {
                "subject_id": subject_id,
                "bank_id": bank_id,
                "difficulty": difficulty,
                "unit": unit
            }
        }
    except Exception as e:
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=f"Search error: {str(e)}")

@app.get("/api/search/papers", tags=["Search"])
def search_papers(
    request: Request,
    q: str = "",
    subject_id: Optional[int] = None,
    limit: int = 50
):
    """Search for question papers"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        allowed = _allowed_subject_ids(connection, user)
        if allowed is not None and not allowed:
            return {
                "success": True,
                "count": 0,
                "results": [],
                "query": q,
                "filters": {"subject_id": subject_id}
            }
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        query = "SELECT id, title, subject_id, exam_type, exam_date, total_marks, generated_at FROM question_papers WHERE 1=1"
        params = []
        
        if q:
            query += f" AND (title LIKE {placeholder} OR exam_type LIKE {placeholder})"
            search_term = f"%{q}%"
            params.extend([search_term, search_term])
        
        if subject_id:
            query += f" AND subject_id = {placeholder}"
            params.append(subject_id)
        
        if allowed is not None:
            id_placeholders = ','.join([placeholder] * len(allowed))
            query += f" AND subject_id IN ({id_placeholders})"
            params.extend(list(allowed))
        
        query += f" ORDER BY generated_at DESC LIMIT {placeholder}"
        params.append(limit)
        
        cursor.execute(query, params)
        results = cursor.fetchall()
        cursor.close()
        connection.close()
        
        return {
            "success": True,
            "count": len(results),
            "results": [dict(r) for r in results],
            "query": q,
            "filters": {"subject_id": subject_id}
        }
    except Exception as e:
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=f"Search error: {str(e)}")

@app.get("/api/search/subjects", tags=["Search"])
def search_subjects(request: Request, q: str = "", limit: int = 50):
    """Search for subjects"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        allowed = _allowed_subject_ids(connection, user)
        if allowed is not None and not allowed:
            return {
                "success": True,
                "count": 0,
                "results": [],
                "query": q
            }
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        query = "SELECT id, subject_id, name FROM subjects WHERE 1=1"
        params = []
        
        if q:
            query += f" AND (name LIKE {placeholder} OR subject_id LIKE {placeholder})"
            search_term = f"%{q}%"
            params.extend([search_term, search_term])
        
        if allowed is not None:
            id_placeholders = ','.join([placeholder] * len(allowed))
            query += f" AND id IN ({id_placeholders})"
            params.extend(list(allowed))
        
        query += f" ORDER BY name LIMIT {placeholder}"
        params.append(limit)
        
        cursor.execute(query, params)
        results = cursor.fetchall()
        cursor.close()
        connection.close()
        
        return {
            "success": True,
            "count": len(results),
            "results": [dict(r) for r in results],
            "query": q
        }
    except Exception as e:
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=f"Search error: {str(e)}")

@app.get("/api/question-images/{image_id}", tags=["Question Images"])
def get_question_image_by_id(image_id: int, request: Request):
    """Get a stored question image (by image id) so previews can embed the exact image."""
    try:
        from services.image_service import ImageService
        from fastapi.responses import StreamingResponse
        import io

        stored = ImageService.get_image_by_id(image_id)
        if not stored or not stored.get('image_blob'):
            raise HTTPException(status_code=404, detail="Image not found")

        image_subject_id = stored.get('subject_id')
        if image_subject_id is not None:
            user = _current_user(request)
            connection = get_db_connection()
            if not connection:
                raise HTTPException(status_code=500, detail="Database connection failed")
            try:
                _assert_subject_id_access(connection, user, int(image_subject_id))
            finally:
                connection.close()

        mime = stored.get('mime_type') or 'image/png'
        media_type = mime if str(mime).startswith('image/') else 'image/png'
        return StreamingResponse(
            io.BytesIO(stored['image_blob']),
            media_type=media_type,
            headers={
                "Content-Disposition": f"inline; filename=question_image_{image_id}.png"
            }
        )
    except HTTPException:
        raise
    except Exception as e:
        print(f"Error fetching image by id {image_id}: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Error fetching image: {str(e)}")


@app.get("/api/questions/{question_id}/image", tags=["Questions"])
def get_question_image(question_id: int, request: Request):
    """Get image for a specific question for preview"""
    try:
        from services.image_integration import get_image_for_question
        from fastapi.responses import StreamingResponse
        import io
        
        connection = get_db_connection()
        if not connection:
            raise HTTPException(status_code=500, detail="Database connection failed")
        
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        # Fetch question content and associated image
        cursor.execute(f"SELECT content, image_id, subject_id FROM questions WHERE id = {placeholder}", (question_id,))
        question = cursor.fetchone()
        
        if not question:
            cursor.close()
            connection.close()
            raise HTTPException(status_code=404, detail="Question not found")
        
        question_subject_id = question.get('subject_id') if isinstance(question, dict) else None
        try:
            if question_subject_id is not None:
                _assert_subject_id_access(connection, _current_user(request), int(question_subject_id))
        finally:
            cursor.close()
            connection.close()
        
        image_blob = None
        media_type = "image/png"
        # Prefer the exact stored image used to generate this question
        stored_image_id = question.get('image_id') if isinstance(question, dict) else None
        if stored_image_id:
            try:
                from services.image_service import ImageService
                stored = ImageService.get_image_by_id(int(stored_image_id))
                if stored and stored.get('image_blob'):
                    image_blob = stored['image_blob']
                    mime = stored.get('mime_type') or 'image/png'
                    media_type = mime if str(mime).startswith('image/') else 'image/png'
            except Exception as exc:
                print(f"Error loading stored image {stored_image_id}: {exc}")
        
        # Fallback: search for a relevant image based on question content
        if not image_blob:
            image_data = get_image_for_question(question['content'], set(), trace_label=f"preview_q{question_id}")
            if image_data and image_data.get('image_blob'):
                image_blob = image_data['image_blob']
        
        if not image_blob:
            raise HTTPException(status_code=404, detail="No image found for this question")
        
        # Return image as blob
        return StreamingResponse(
            io.BytesIO(image_blob),
            media_type=media_type,
            headers={
                "Content-Disposition": f"inline; filename=question_{question_id}_image.png"
            }
        )
    except HTTPException:
        raise
    except Exception as e:
        print(f"Error fetching question image: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Error fetching image: {str(e)}")

# ==================== BLUEPRINT ENDPOINTS WITH PARTS ====================

def sync_blueprint_data(connection, blueprint_id, parts_data):
    """Sync blueprint parts and update total counts/marks."""
    cursor = get_cursor(connection)
    placeholder = get_placeholder()
    
    try:
        # Delete existing parts
        cursor.execute(f"DELETE FROM blueprint_parts WHERE blueprint_id = {placeholder}", (blueprint_id,))
        
        total_questions = 0
        total_marks = 0
        
        for i, part in enumerate(parts_data):
            # Support both frontend format (part_name, num_questions, marks_per_question)
            # and DEFAULT_BLUEPRINT_STRUCTURE format (name, count, marks_per_question)
            part_name = part.get('part_name') or part.get('name')
            instructions = part.get('instructions') or part.get('instruction') or "Answer all questions."
            num_questions = part.get('num_questions') or part.get('count') or 0
            marks_per = part.get('marks_per_question') or 0
            difficulty = part.get('difficulty') or "medium"
            effective_count = _effective_count_from_instruction(str(instructions), int(num_questions))
            
            cursor.execute(
                f"""INSERT INTO blueprint_parts 
                   (blueprint_id, part_name, instructions, num_questions, marks_per_question, difficulty, part_order)
                   VALUES ({placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder})""",
                (blueprint_id, part_name, instructions, num_questions, marks_per, difficulty, i)
            )
            
            total_questions += int(num_questions)
            total_marks += (effective_count * float(marks_per))
            
        # Update blueprint totals
        cursor.execute(
            f"UPDATE blueprints SET total_questions = {placeholder}, total_marks = {placeholder}, parts_config = {placeholder} WHERE id = {placeholder}",
            (total_questions, total_marks, json.dumps(parts_data), blueprint_id)
        )
        connection.commit()
    finally:
        cursor.close()

def ensure_default_blueprint():
    """Ensure the database has the blueprints table and at least one default blueprint."""
    init_database()

    connection = get_db_connection()
    if not connection:
        print("Error ensuring default blueprint: database connection failed")
        return None

    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()

        cursor.execute("SELECT COUNT(*) as count FROM blueprints")
        count_row = cursor.fetchone()
        count = count_row['count'] if isinstance(count_row, dict) or hasattr(count_row, '__getitem__') else 0

        if count == 0:
            blueprints_dir = UPLOAD_DIR / "blueprints"
            blueprints_dir.mkdir(exist_ok=True)

            file_name = "default_blueprint.json"
            file_path = blueprints_dir / file_name

            with open(file_path, 'w', encoding='utf-8') as f:
                json.dump(DEFAULT_BLUEPRINT_STRUCTURE, f, indent=2)

            cursor.execute(
                f"""INSERT INTO blueprints (name, description, file_name, file_path) 
                   VALUES ({placeholder}, {placeholder}, {placeholder}, {placeholder})""",
                (
                    DEFAULT_BLUEPRINT_STRUCTURE['name'],
                    DEFAULT_BLUEPRINT_STRUCTURE['description'],
                    file_name,
                    str(file_path)
                )
            )
            connection.commit()
            blueprint_id = cursor.lastrowid
            
            # Sync parts
            sync_blueprint_data(connection, blueprint_id, DEFAULT_BLUEPRINT_STRUCTURE['parts'])

        cursor.close()
        connection.close()
        return True
    except Exception as e:
        if "no such table" in str(e).lower():
            print("Blueprints table missing. Re-initializing database...")
            init_database()
            return ensure_default_blueprint()
        print(f"Error ensuring default blueprint: {e}")
        if connection:
            connection.close()
        return None


def ensure_default_blueprint_exists():
    """Backward-compatible wrapper used by other endpoints."""
    return ensure_default_blueprint()

@app.post("/api/blueprints", tags=["Blueprints"], response_model=BlueprintResponse)
def create_blueprint(blueprint: BlueprintCreate, request: Request):
    """Create a new blueprint from JSON structure (admin only)"""
    _require_roles(request, "admin")
    
    print("=" * 60)
    print("📥 RECEIVED BLUEPRINT DATA:")
    print(f"Name: {blueprint.name}")
    print(f"Description: {blueprint.description}")
    print(f"Parts config: {blueprint.parts_config}")
    print(f"Number of parts: {len(blueprint.parts_config)}")
    for i, part in enumerate(blueprint.parts_config):
        print(f"  Part {i+1}: {part.part_name} - {part.difficulty} - {part.num_questions}q × {part.marks_per_question}m")
    print("=" * 60)
    
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        blueprints_dir = UPLOAD_DIR / "blueprints"
        blueprints_dir.mkdir(exist_ok=True)
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_name = "".join(c if c.isalnum() or c in (' ', '-', '_') else '_' for c in blueprint.name)
        safe_name = safe_name.replace(' ', '_')[:50]
        file_name = f"{safe_name}_{timestamp}.json"
        file_path = str(blueprints_dir / file_name)
        
        total_questions = sum(part.num_questions for part in blueprint.parts_config)
        total_marks = sum(part.num_questions * part.marks_per_question for part in blueprint.parts_config)
        
        json_data = {
            "name": blueprint.name,
            "description": blueprint.description or "",
            "total_marks": total_marks,
            "total_questions": total_questions,
            "parts": [
                {
                    "part_name": part.part_name,
                    "instructions": part.instructions or "Answer all questions",
                    "num_questions": part.num_questions,
                    "marks_per_question": part.marks_per_question,
                    "difficulty": part.difficulty
                }
                for part in blueprint.parts_config
            ]
        }
        
        with open(file_path, "w", encoding='utf-8') as f:
            json.dump(json_data, f, indent=2, ensure_ascii=False)
        
        print(f"✅ Saved blueprint file to: {file_path}")
        
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        cursor.execute(
            f"""INSERT INTO blueprints (name, description, file_name, file_path, total_questions, total_marks) 
               VALUES ({placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder})""",
            (blueprint.name, blueprint.description, file_name, file_path, total_questions, total_marks)
        )
        connection.commit()
        blueprint_id = cursor.lastrowid
        
        print(f"✅ Created blueprint with ID: {blueprint_id}")
        
        # Insert blueprint parts
        for i, part in enumerate(blueprint.parts_config):
            cursor.execute(
                f"""INSERT INTO blueprint_parts 
                   (blueprint_id, part_order, part_name, instructions, num_questions, marks_per_question, difficulty)
                   VALUES ({placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder})""",
                (blueprint_id, i + 1, part.part_name, part.instructions or "Answer all questions", 
                 part.num_questions, part.marks_per_question, part.difficulty)
            )
            print(f"  ✅ Inserted part: {part.part_name} (order: {i + 1})")
        
        connection.commit()
        print(f"✅ Inserted {len(blueprint.parts_config)} parts for blueprint {blueprint_id}")
        
        cursor.execute(f"SELECT * FROM blueprints WHERE id = {placeholder}", (blueprint_id,))
        result = cursor.fetchone()
        
        cursor.close()
        connection.close()
        
        return dict(result)
        
    except Exception as e:
        print(f"❌ ERROR: {str(e)}")
        import traceback
        traceback.print_exc()
        
        if 'file_path' in locals() and os.path.exists(file_path):
            os.remove(file_path)
        if connection:
            connection.rollback()
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error creating blueprint: {str(e)}")

@app.put("/api/blueprints/{blueprint_id}", tags=["Blueprints"], response_model=BlueprintResponse)
def update_blueprint(blueprint_id: int, blueprint: BlueprintCreate, request: Request):
    """Update an existing blueprint (name, description, and parts structure). Admin only."""
    _require_roles(request, "admin")
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")

    try:
        placeholder = get_placeholder()
        cursor = get_cursor(connection)

        cursor.execute(f"SELECT id, file_path FROM blueprints WHERE id = {placeholder}", (blueprint_id,))
        existing = cursor.fetchone()
        if not existing:
            connection.close()
            raise HTTPException(status_code=404, detail="Blueprint not found")

        total_questions = sum(part.num_questions for part in blueprint.parts_config)
        total_marks = sum(part.num_questions * part.marks_per_question for part in blueprint.parts_config)

        json_data = {
            "name": blueprint.name,
            "description": blueprint.description or "",
            "total_marks": total_marks,
            "total_questions": total_questions,
            "parts": [
                {
                    "part_name": part.part_name,
                    "instructions": part.instructions or "Answer all questions",
                    "num_questions": part.num_questions,
                    "marks_per_question": part.marks_per_question,
                    "difficulty": part.difficulty
                }
                for part in blueprint.parts_config
            ]
        }

        # Reuse the existing JSON file path when available to avoid orphaned files
        file_path = None
        existing_path = existing.get("file_path") if existing else None
        if existing_path:
            try:
                with open(existing_path, "w", encoding='utf-8') as f:
                    json.dump(json_data, f, indent=2, ensure_ascii=False)
                file_path = existing_path
            except Exception:
                file_path = None

        if not file_path:
            blueprints_dir = UPLOAD_DIR / "blueprints"
            blueprints_dir.mkdir(exist_ok=True)
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            safe_name = "".join(c if c.isalnum() or c in (' ', '-', '_') else '_' for c in blueprint.name)
            safe_name = safe_name.replace(' ', '_')[:50]
            file_name = f"{safe_name}_{timestamp}.json"
            file_path = str(blueprints_dir / file_name)
            with open(file_path, "w", encoding='utf-8') as f:
                json.dump(json_data, f, indent=2, ensure_ascii=False)

        cursor.execute(
            f"""UPDATE blueprints 
               SET name = {placeholder}, description = {placeholder}, file_path = {placeholder}, 
                   total_questions = {placeholder}, total_marks = {placeholder}, updated_at = CURRENT_TIMESTAMP
               WHERE id = {placeholder}""",
            (blueprint.name, blueprint.description, file_path, total_questions, total_marks, blueprint_id)
        )

        cursor.execute(f"DELETE FROM blueprint_parts WHERE blueprint_id = {placeholder}", (blueprint_id,))
        for i, part in enumerate(blueprint.parts_config):
            cursor.execute(
                f"""INSERT INTO blueprint_parts 
                   (blueprint_id, part_order, part_name, instructions, num_questions, marks_per_question, difficulty)
                   VALUES ({placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder})""",
                (blueprint_id, i + 1, part.part_name, part.instructions or "Answer all questions",
                 part.num_questions, part.marks_per_question, part.difficulty)
            )

        connection.commit()

        cursor.execute(f"SELECT * FROM blueprints WHERE id = {placeholder}", (blueprint_id,))
        result = cursor.fetchone()

        cursor.close()
        connection.close()

        print(f"✅ Updated blueprint ID: {blueprint_id}")
        return dict(result)

    except Exception as e:
        print(f"❌ ERROR updating blueprint: {str(e)}")
        import traceback
        traceback.print_exc()
        if connection:
            connection.rollback()
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error updating blueprint: {str(e)}")


@app.get("/api/blueprints", tags=["Blueprints"], response_model=List[BlueprintResponse])
def get_blueprints():
    """Get all blueprints"""
    ensure_default_blueprint()
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        cursor.execute("SELECT * FROM blueprints ORDER BY created_at DESC")
        blueprints = cursor.fetchall()
        
        cursor.close()
        connection.close()
        
        return [dict(p) for p in blueprints]
    except Exception as e:
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error fetching blueprints: {str(e)}")



@app.get("/api/blueprints/{blueprint_id}", tags=["Blueprints"])
def get_blueprint(blueprint_id: int):
    """Get a specific blueprint with its parts"""
    try:
        # Verify blueprint exists
        BlueprintGuard.verify_existence(blueprint_id)
        
        # Get blueprint with parts
        blueprint_dict = BlueprintRepository.get_with_parts(blueprint_id)
        
        if not blueprint_dict:
            raise HTTPException(status_code=404, detail="Blueprint not found")
        
        return blueprint_dict
        
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error fetching blueprint: {str(e)}")



@app.delete("/api/blueprints/{blueprint_id}", tags=["Blueprints"], status_code=status.HTTP_204_NO_CONTENT)
def delete_blueprint(blueprint_id: int, request: Request):
    """Delete a blueprint (admin only)"""
    _require_roles(request, "admin")
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        cursor.execute(f"SELECT file_path FROM blueprints WHERE id = {placeholder}", (blueprint_id,))
        blueprint = cursor.fetchone()
        
        if not blueprint:
            cursor.close()
            connection.close()
            raise HTTPException(status_code=404, detail="Blueprint not found")
        
        blueprint_dict = dict(blueprint)
        
        if blueprint_dict.get('file_path') and os.path.exists(blueprint_dict['file_path']):
            os.remove(blueprint_dict['file_path'])
        
        cursor.execute(f"DELETE FROM blueprints WHERE id = {placeholder}", (blueprint_id,))
        connection.commit()
        
        cursor.close()
        connection.close()
        
        return None
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.rollback()
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error deleting blueprint: {str(e)}")

# ==================== DASHBOARD ENDPOINTS ====================

@app.get("/api/dashboard/stats", tags=["Dashboard"])
def get_dashboard_stats(request: Request):
    """Get dashboard statistics scoped to the caller's access"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        allowed = _allowed_subject_ids(connection, user)

        if allowed is None:
            cursor.execute("SELECT COUNT(*) as count FROM subjects")
            subjects_count = cursor.fetchone()['count']

            cursor.execute("SELECT COUNT(*) as count FROM questions")
            questions_count = cursor.fetchone()['count']

            cursor.execute("SELECT COUNT(*) as count FROM question_papers")
            papers_count = cursor.fetchone()['count']
        elif allowed:
            id_placeholders = ','.join([placeholder] * len(allowed))
            id_list = list(allowed)

            cursor.execute(f"SELECT COUNT(*) as count FROM subjects WHERE id IN ({id_placeholders})", tuple(id_list))
            subjects_count = cursor.fetchone()['count']

            cursor.execute(f"SELECT COUNT(*) as count FROM questions WHERE subject_id IN ({id_placeholders})", tuple(id_list))
            questions_count = cursor.fetchone()['count']

            cursor.execute(f"SELECT COUNT(*) as count FROM question_papers WHERE subject_id IN ({id_placeholders})", tuple(id_list))
            papers_count = cursor.fetchone()['count']
        else:
            subjects_count = questions_count = papers_count = 0

        cursor.execute("SELECT COUNT(*) as count FROM blueprints")
        blueprints_count = cursor.fetchone()['count']
        
        cursor.close()
        connection.close()
        
        return {
            "subjects": subjects_count,
            "questions": questions_count,
            "blueprints": blueprints_count,
            "papers": papers_count
        }
    except Exception as e:
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error fetching stats: {str(e)}")

@app.get("/api/dashboard/recent-activity", tags=["Dashboard"])
def get_recent_activity(request: Request):
    """Get recent activity from multiple sources (scoped to the caller)"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        
        activities = []
        
        cursor.execute("""
            SELECT qb.name, s.name as subject_name, qb.created_at, 'question_bank' as type
            FROM question_banks qb
            JOIN subjects s ON qb.subject_id = s.id
            ORDER BY qb.created_at DESC
            LIMIT 5
        """)
        qbanks = cursor.fetchall()
        for qb in qbanks:
            activities.append({
                "action": f"Created question bank '{qb['name']}'",
                "subject": qb['subject_name'],
                "time": qb['created_at'],
                "type": "question_bank"
            })
        
        cursor.execute("""
            SELECT name, created_at
            FROM subjects
            ORDER BY created_at DESC
            LIMIT 3
        """)
        subjects = cursor.fetchall()
        for subj in subjects:
            activities.append({
                "action": "Added new subject",
                "subject": subj['name'],
                "time": subj['created_at'],
                "type": "subject"
            })
        
        cursor.execute("""
            SELECT name, created_at
            FROM blueprints
            ORDER BY created_at DESC
            LIMIT 3
        """)
        blueprints = cursor.fetchall()
        for bp in blueprints:
            activities.append({
                "action": f"Created blueprint '{bp['name']}'",
                "subject": "N/A",
                "time": bp['created_at'],
                "type": "blueprint"
            })
        
        cursor.execute("""
            SELECT qp.title, s.name as subject_name, qp.generated_at
            FROM question_papers qp
            JOIN subjects s ON qp.subject_id = s.id
            ORDER BY qp.generated_at DESC
            LIMIT 3
        """)
        papers = cursor.fetchall()
        for paper in papers:
            activities.append({
                "action": "Generated question paper",
                "subject": paper['subject_name'],
                "time": paper['generated_at'],
                "type": "paper"
            })
        
        activities.sort(key=lambda x: x['time'], reverse=True)
        activities = activities[:10]

        if user["role"] != "admin":
            visible_names = {(r["name"] or "").strip().lower() for r in _visible_subject_rows(connection, user)}
            activities = [
                a for a in activities
                if a.get("type") != "blueprint" and (a.get("subject") or "").strip().lower() in visible_names
            ]
        
        cursor.close()
        connection.close()
        
        return activities
    except Exception as e:
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error fetching recent activity: {str(e)}")

# ==================== QUESTION PAPER ENDPOINTS ====================

@app.post("/api/question-papers/generate", tags=["Question Papers"])
def generate_question_paper_endpoint(
    http_request: Request,
    title: str = Form(...),
    subject_id: int = Form(...),
    question_bank_id: int = Form(...),
    blueprint_id: Optional[int] = Form(None),
    exam_type: Optional[str] = Form("Regular"),
    exam_date: Optional[str] = Form(None),
    duration: Optional[str] = Form("3"),
    file_format: str = Form("pdf"),
    need_image: Optional[str] = Form("no"),
    image_sources: Optional[str] = Form(None)
):
    """Generate actual question paper document from question bank"""
    user = _current_user(http_request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        _assert_subject_id_access(connection, user, subject_id)
        
        if blueprint_id is None:
            ensure_default_blueprint_exists()
            cursor.execute("SELECT id FROM blueprints ORDER BY created_at ASC LIMIT 1")
            first_bp = cursor.fetchone()
            if first_bp:
                blueprint_id = first_bp['id']
        
        cursor.execute(f"SELECT id, name, subject_id FROM subjects WHERE id = {placeholder}", (subject_id,))
        subject = cursor.fetchone()
        if not subject:
            raise HTTPException(status_code=404, detail="Subject not found")
        
        cursor.execute(f"SELECT id, name FROM question_banks WHERE id = {placeholder}", (question_bank_id,))
        qbank = cursor.fetchone()
        if not qbank:
            raise HTTPException(status_code=404, detail="Question bank not found")
        
        cursor.execute(f"SELECT id, name, file_path FROM blueprints WHERE id = {placeholder}", (blueprint_id,))
        blueprint = cursor.fetchone()
        if not blueprint:
            raise HTTPException(status_code=404, detail="Blueprint not found")
        
        if not blueprint['file_path']:
            raise HTTPException(
                status_code=400, 
                detail=f"Blueprint '{blueprint['name']}' (ID: {blueprint['id']}) has no file attached. Please delete and recreate it with a JSON file."
            )
        
        if not os.path.exists(blueprint['file_path']):
            raise HTTPException(
                status_code=404, 
                detail=f"Blueprint file not found: {blueprint['file_path']}. Please re-upload the blueprint."
            )
        
        file_ext = blueprint['file_path'].lower()
        if not (file_ext.endswith('.json') or file_ext.endswith('.docx') or file_ext.endswith('.doc')):
            raise HTTPException(status_code=400, detail="Blueprint must be a JSON or DOCX file.")
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_title = "".join(c for c in title if c.isalnum() or c in (' ', '-', '_')).strip()
        safe_title = safe_title.replace(' ', '_')
        filename = f"{subject['subject_id']}_{safe_title}_{timestamp}.{file_format}"
        output_path = PAPERS_DIR / filename
        
        # Parse need_image and image_sources
        is_need_image = (need_image or "no").strip().lower() in ("yes", "true", "1")
        sources_list = []
        if image_sources:
            if image_sources.startswith("["):
                try:
                    sources_list = json.loads(image_sources)
                except:
                    sources_list = [s.strip() for s in image_sources.split(",") if s.strip()]
            else:
                sources_list = [s.strip() for s in image_sources.split(",") if s.strip()]

        output_path_str, questions_by_part = generate_question_paper(
            cursor=cursor,
            title=title,
            subject_id=subject_id,
            subject_name=subject['name'],
            question_bank_id=question_bank_id,
            blueprint_path=blueprint['file_path'],
            exam_type=exam_type,
            exam_date=exam_date,
            duration=duration,
            file_format=file_format,
            output_path=str(output_path),
            need_image=is_need_image,
            image_sources=sources_list
        )
        
        parsed_date = None
        if exam_date:
            try:
                from datetime import datetime as dt
                parsed_date = dt.strptime(exam_date, "%Y-%m-%d").date()
            except:
                pass
        
        total_marks = 0
        try:
            if blueprint['file_path'].lower().endswith('.json'):
                with open(blueprint['file_path'], 'r', encoding='utf-8') as f:
                    content = f.read()
                    if not content.strip():
                        raise HTTPException(status_code=400, detail="Blueprint file is empty")
                    bp_data = json.loads(content)
                    parts = bp_data.get('parts', [])
                    for part in parts:
                        instructions = part.get('instructions') or part.get('instruction') or "Answer all questions"
                        configured_count = int(part.get('num_questions') or part.get('count') or 0)
                        marks_per = float(part.get('marks_per_question') or 0)
                        effective_count = _effective_count_from_instruction(str(instructions), configured_count)
                        total_marks += effective_count * marks_per
        except (UnicodeDecodeError, json.JSONDecodeError):
            total_marks = 100

        if total_marks <= 0:
            total_marks = 100
        
        questions_data_json = json.dumps(questions_by_part)
        
        query = f"""
            INSERT INTO question_papers 
            (title, subject_id, blueprint_id, exam_type, exam_date, total_marks, file_format, file_path, questions_data)
            VALUES ({placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder})
        """
        cursor.execute(query, (title, subject_id, blueprint_id, exam_type, parsed_date, total_marks, file_format, str(output_path), questions_data_json))
        connection.commit()
        
        paper_id = cursor.lastrowid
        
        cursor.execute(f"""
            SELECT id, title, subject_id, blueprint_id, exam_type, exam_date, 
                   total_marks, file_format, file_path, generated_at
            FROM question_papers WHERE id = {placeholder}
        """, (paper_id,))
        paper = cursor.fetchone()
        
        cursor.close()
        connection.close()
        
        return dict(paper)
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.rollback()
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error generating question paper: {str(e)}")

@app.post("/api/question-papers/generate-from-data", tags=["Question Papers"], response_model=QuestionPaperResponse)
def generate_question_paper_from_data(request: dict, http_request: Request):
    """Generate PDF/DOCX from paper data sent by frontend"""
    user = _current_user(http_request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        # Extract data from request
        title = request.get('title')
        subject_id = request.get('subject_id')
        blueprint_id = request.get('blueprint_id')
        exam_type = request.get('exam_type', 'Regular')
        exam_date = request.get('exam_date')
        exam_duration = request.get('exam_duration', '3')
        total_marks = 0
        file_format = request.get('file_format', 'pdf')
        paper_data = request.get('paper_data')  # Contains parts and questions
        
        # Validate subject
        cursor.execute(f"SELECT id, subject_id, name, course_outcome_file FROM subjects WHERE id = {placeholder}", (subject_id,))
        subject = cursor.fetchone()
        if not subject:
            raise HTTPException(status_code=404, detail="Subject not found")
        _assert_subject_access(user, dict(subject), connection=connection)

        course_outcome_file = None
        try:
            course_outcome_file = subject['course_outcome_file']
        except Exception:
            course_outcome_file = None
        
        # Create output file
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_title = "".join(c for c in title if c.isalnum() or c in (' ', '-', '_')).strip()
        safe_title = safe_title.replace(' ', '_')
        filename = f"{subject['subject_id']}_{safe_title}_{timestamp}.{file_format}"
        output_path = PAPERS_DIR / filename
        
        # Convert paper_data to questions_by_part format
        questions_by_part = {}
        for part in paper_data.get('parts', []):
            questions_by_part[part['part_name']] = part['questions']
        
        # Create blueprint dict from paper_data
        blueprint = {
            'name': title,
            'total_marks': 0,
            'parts': [
                {
                    'part_name': part['part_name'],
                    'instructions': part.get('instructions', 'Answer all questions'),
                    'num_questions': len(part['questions']),
                    'marks_per_question': part.get('marks_per_question', 2),
                    'difficulty': part.get('difficulty', 'medium')
                }
                for part in paper_data.get('parts', [])
            ]
        }

        # Calculate effective total marks based on instructions (e.g., "Answer any two")
        for part in blueprint['parts']:
            instructions = part.get('instructions') or part.get('instruction') or "Answer all questions"
            configured_count = int(part.get('num_questions') or part.get('count') or 0)
            marks_per = float(part.get('marks_per_question') or 0)
            effective_count = _effective_count_from_instruction(str(instructions), configured_count)
            total_marks += effective_count * marks_per

        if total_marks <= 0:
            total_marks = request.get('total_marks', 100)

        blueprint['total_marks'] = total_marks
        
        need_image = request.get('need_image', False)
        if isinstance(need_image, str):
            need_image = need_image.strip().lower() in ('yes', 'true', '1')
        image_sources = request.get('image_sources', [])

        # Generate the file
        if file_format == 'pdf':
            from services.paper_generator import generate_pdf_paper
            generate_pdf_paper(
                title=title,
                subject_name=subject['name'],
                exam_type=exam_type,
                exam_date=exam_date or '',
                total_marks=total_marks,
                duration=exam_duration,
                blueprint=blueprint,
                questions_by_part=questions_by_part,
                output_path=str(output_path),
                course_outcome_file=course_outcome_file,
                subject_code=subject.get('subject_id'),
                need_image=need_image,
                image_sources=image_sources,
            )
        else:  # docx
            from services.paper_generator import generate_docx_paper
            generate_docx_paper(
                title=title,
                subject_name=subject['name'],
                exam_type=exam_type,
                exam_date=exam_date or '',
                total_marks=total_marks,
                duration=exam_duration,
                blueprint=blueprint,
                questions_by_part=questions_by_part,
                output_path=str(output_path),
                course_outcome_file=course_outcome_file,
                need_image=need_image,
                image_sources=image_sources,
            )
        
        # Parse exam date
        parsed_date = None
        if exam_date:
            try:
                from datetime import datetime as dt
                parsed_date = dt.strptime(exam_date, "%Y-%m-%d").date()
            except:
                pass
        
        # Save to database
        questions_data_json = json.dumps(questions_by_part)
        
        query = f"""
            INSERT INTO question_papers 
            (title, subject_id, blueprint_id, exam_type, exam_date, total_marks, file_format, file_path, questions_data)
            VALUES ({placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder})
        """
        
        cursor.execute(
            query,
            (title, subject_id, blueprint_id, exam_type, parsed_date, total_marks, file_format, str(output_path), questions_data_json)
        )
        
        connection.commit()
        
        # Get the inserted paper
        paper_id = cursor.lastrowid if get_db_type() != 'postgresql' else cursor.fetchone()[0]
        cursor.execute(f"SELECT * FROM question_papers WHERE id = {placeholder}", (paper_id,))
        paper = cursor.fetchone()
        
        cursor.close()
        connection.close()
        
        return dict(paper)
        
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.rollback()
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error generating paper: {str(e)}")

@app.post("/api/question-papers", tags=["Question Papers"], response_model=QuestionPaperResponse)
def create_question_paper(
    request: Request,
    title: str = Form(...),
    subject_id: int = Form(...),
    blueprint_id: Optional[int] = Form(None),
    exam_type: Optional[str] = Form(None),
    exam_date: Optional[str] = Form(None),
    total_marks: Optional[float] = Form(None),
    file_format: Optional[str] = Form("txt"),
    paper_content: UploadFile = File(...)
):
    """Save a pre-generated question paper to database (legacy endpoint)"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        _assert_subject_id_access(connection, user, subject_id)
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        # Validate subject exists
        cursor.execute(f"SELECT id, subject_id, name FROM subjects WHERE id = {placeholder}", (subject_id,))
        subject = cursor.fetchone()
        if not subject:
            cursor.close()
            connection.close()
            raise HTTPException(status_code=404, detail="Subject not found")
        
        # Save the uploaded file
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_title = "".join(c for c in title if c.isalnum() or c in (' ', '-', '_')).strip()
        safe_title = safe_title.replace(' ', '_')
        
        # Determine file extension
        content_type = paper_content.content_type or ''
        if 'pdf' in content_type or file_format == 'pdf':
            ext = 'pdf'
        elif 'word' in content_type or 'officedocument' in content_type or file_format in ['docx', 'doc']:
            ext = 'docx'
        else:
            ext = 'txt'
        
        filename = f"{subject['subject_id']}_{safe_title}_{timestamp}.{ext}"
        file_path = PAPERS_DIR / filename
        
        # Save file
        content = paper_content.file.read()
        with open(file_path, 'wb') as f:
            f.write(content)
        
        # Parse exam date
        parsed_date = None
        if exam_date:
            try:
                from datetime import datetime as dt
                parsed_date = dt.strptime(exam_date, "%Y-%m-%d").date()
            except:
                pass
        
        # Insert into database
        query = f"""
            INSERT INTO question_papers 
            (title, subject_id, blueprint_id, exam_type, exam_date, total_marks, file_format, file_path, questions_data)
            VALUES ({placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder})
        """
        
        cursor.execute(
            query,
            (title, subject_id, blueprint_id, exam_type, parsed_date, total_marks or 100, ext, str(file_path), None)
        )
        
        connection.commit()
        
        # Get the inserted paper
        paper_id = cursor.lastrowid if get_db_type() != 'postgresql' else cursor.fetchone()[0]
        cursor.execute(f"SELECT * FROM question_papers WHERE id = {placeholder}", (paper_id,))
        paper = cursor.fetchone()
        
        cursor.close()
        connection.close()
        
        return dict(paper)
        
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.rollback()
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error saving question paper: {str(e)}")

@app.get("/api/question-papers", tags=["Question Papers"], response_model=List[QuestionPaperResponse])
def get_all_question_papers(request: Request):
    """Get question papers visible to the caller"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        allowed = _allowed_subject_ids(connection, user)
        cursor = get_cursor(connection)
        if allowed is None:
            cursor.execute("""
                SELECT qp.*, s.name as subject_name,
                       EXISTS(SELECT 1 FROM answer_scripts WHERE question_paper_id = qp.id) as has_answer_script
                FROM question_papers qp
                LEFT JOIN subjects s ON qp.subject_id = s.id
                ORDER BY qp.generated_at DESC
            """)
        elif allowed:
            id_placeholders = ','.join(['%s'] * len(allowed))
            cursor.execute(f"""
                SELECT qp.*, s.name as subject_name,
                       EXISTS(SELECT 1 FROM answer_scripts WHERE question_paper_id = qp.id) as has_answer_script
                FROM question_papers qp
                LEFT JOIN subjects s ON qp.subject_id = s.id
                WHERE qp.subject_id IN ({id_placeholders})
                ORDER BY qp.generated_at DESC
            """, tuple(allowed))
        else:
            cursor.execute("SELECT 1 FROM dual WHERE 1=0")
        papers = cursor.fetchall()
        cursor.close()
        connection.close()
        return [dict(p) for p in papers]
    except Exception as e:
        print(f"DEBUG: Error in get_all_question_papers: {str(e)}")
        if connection:
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error fetching question papers: {str(e)}")

@app.get("/api/question-papers/{paper_id}/download", tags=["Question Papers"])
def download_question_paper(paper_id: int, request: Request):
    """Download the generated question paper file"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        _assert_paper_access(connection, user, paper_id)
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        cursor.execute(f"SELECT file_path, title, file_format FROM question_papers WHERE id = {placeholder}", (paper_id,))
        paper = cursor.fetchone()
        
        cursor.close()
        connection.close()
        
        if not paper:
            raise HTTPException(status_code=404, detail="Question paper not found")
        
        if not paper['file_path'] or not os.path.exists(paper['file_path']):
            raise HTTPException(status_code=404, detail="Question paper file not found")
        
        # Determine media type based on file format
        file_format = paper['file_format'] or 'pdf'
        if file_format == 'pdf':
            media_type = "application/pdf"
        elif file_format in ['docx', 'doc']:
            media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        else:
            media_type = "text/plain"
        
        filename = os.path.basename(paper['file_path'])
        
        return FileResponse(
            path=paper['file_path'],
            media_type=media_type,
            filename=filename
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error downloading paper: {str(e)}")

@app.delete("/api/question-papers/{paper_id}", tags=["Question Papers"])
def delete_question_paper(paper_id: int, request: Request):
    """Delete a question paper"""
    user = _current_user(request)
    connection = get_db_connection()
    if not connection:
        raise HTTPException(status_code=500, detail="Database connection failed")
    
    try:
        _assert_paper_access(connection, user, paper_id)
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        cursor.execute(f"SELECT file_path FROM question_papers WHERE id = {placeholder}", (paper_id,))
        result = cursor.fetchone()
        
        if not result:
            cursor.close()
            connection.close()
            raise HTTPException(status_code=404, detail="Question paper not found")
        
        file_path = result['file_path'] if result else None
        
        cursor.execute(f"DELETE FROM question_papers WHERE id = {placeholder}", (paper_id,))
        connection.commit()
        
        if file_path and os.path.exists(file_path):
            os.remove(file_path)
        
        cursor.close()
        connection.close()
        
        return {"message": "Question paper deleted successfully"}
    except HTTPException:
        raise
    except Exception as e:
        if connection:
            connection.rollback()
            connection.close()
        raise HTTPException(status_code=500, detail=f"Error deleting question paper: {str(e)}")

# ==================== GRADING & EVALUATION ENDPOINTS ====================

@app.post("/api/answer-scripts/generate/{paper_id}", tags=["Answer Scripts"])
def generate_script(paper_id: int, http_request: Request):
    """Generate answer script for a question paper"""
    user = _current_user(http_request)
    connection = get_db_connection()
    if connection:
        _assert_paper_access(connection, user, paper_id)
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()

        # Fast path: if answer script already exists, return immediately
        cursor.execute(
            f"SELECT id, created_at FROM answer_scripts WHERE question_paper_id = {placeholder} ORDER BY created_at DESC LIMIT 1",
            (paper_id,)
        )
        existing_script = cursor.fetchone()
        if existing_script:
            cursor.close()
            connection.close()
            return {
                "success": True,
                "message": "Answer script already exists",
                "already_exists": True,
                "script_id": existing_script["id"] if isinstance(existing_script, dict) else existing_script[0],
            }
        
        cursor.execute(f"SELECT questions_data FROM question_papers WHERE id = {placeholder}", (paper_id,))
        paper = cursor.fetchone()
        if not paper:
            raise HTTPException(status_code=404, detail="Question paper not found")
        
        if not paper['questions_data']:
            raise HTTPException(
                status_code=400, 
                detail="This paper was generated before the automated grading update and doesn't store question data. Please generate a new question paper to use AI grading."
            )
        
        questions_by_part = json.loads(paper['questions_data'])
        answers = generate_answer_script(questions_by_part)
        
        if not answers:
            raise HTTPException(status_code=500, detail="Failed to generate answers using AI")
        
        answer_data_json = json.dumps(answers)
        
        # Save to DB
        cursor.execute(
            f"INSERT INTO answer_scripts (question_paper_id, answer_data) VALUES ({placeholder}, {placeholder})",
            (paper_id, answer_data_json)
        )
        connection.commit()
        
        cursor.close()
        connection.close()
        return {"success": True, "message": "Answer script generated successfully"}
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        traceback.print_exc()
        if connection: connection.close()
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/answer-scripts/{paper_id}", tags=["Answer Scripts"])
def get_script(paper_id: int, http_request: Request):
    """Get answer script for a paper"""
    user = _current_user(http_request)
    connection = get_db_connection()
    if connection:
        _assert_paper_access(connection, user, paper_id)
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        cursor.execute(f"SELECT * FROM answer_scripts WHERE question_paper_id = {placeholder} ORDER BY created_at DESC LIMIT 1", (paper_id,))
        script = cursor.fetchone()
        cursor.close()
        connection.close()
        if not script:
            raise HTTPException(status_code=404, detail="Answer script not found")
        return dict(script)
    except HTTPException:
        raise
    except Exception as e:
        if connection: connection.close()
        raise HTTPException(status_code=500, detail=str(e))

@app.put("/api/answer-scripts/{paper_id}", tags=["Answer Scripts"])
def update_script(http_request: Request, paper_id: int, request_data: dict = Body(...)):
    """Update answer script for a paper"""
    user = _current_user(http_request)
    connection = get_db_connection()
    if connection:
        _assert_paper_access(connection, user, paper_id)
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        # Ensure question paper exists
        cursor.execute(f"SELECT id FROM question_papers WHERE id = {placeholder}", (paper_id,))
        if not cursor.fetchone():
            raise HTTPException(status_code=404, detail="Question paper not found")
        
        if 'answer_data' not in request_data:
            raise HTTPException(status_code=400, detail=f"Answer data is required. received keys: {list(request_data.keys())}")
            
        answer_data = request_data.get('answer_data')
        if answer_data is None:
            raise HTTPException(status_code=400, detail="Answer data cannot be null")
            
        # Check if it exists
        cursor.execute(f"SELECT id FROM answer_scripts WHERE question_paper_id = {placeholder} ORDER BY created_at DESC LIMIT 1", (paper_id,))
        existing = cursor.fetchone()
        
        if existing:
            # Update existing
            cursor.execute(
                f"UPDATE answer_scripts SET answer_data = {placeholder} WHERE id = {placeholder}",
                (json.dumps(answer_data), existing['id'])
            )
        else:
            # Create new if doesn't exist (though it should)
            cursor.execute(
                f"INSERT INTO answer_scripts (question_paper_id, answer_data) VALUES ({placeholder}, {placeholder})",
                (paper_id, json.dumps(answer_data))
            )
            
        connection.commit()
        cursor.close()
        connection.close()
        return {"success": True, "message": "Answer script updated successfully"}
    except HTTPException:
        raise
    except Exception as e:
        if connection: connection.close()
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/evaluations/evaluate", tags=["Evaluations"])
def evaluate_student(
    http_request: Request,
    paper_id: int = Form(...),
    student_name: str = Form(...),
    register_number: str = Form(...),
    department: str = Form(...),
    student_file: UploadFile = File(...)
):
    """Upload student paper and evaluate using AI"""
    user = _current_user(http_request)
    connection = get_db_connection()
    if connection:
        _assert_paper_access(connection, user, paper_id)
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        # 1. Save uploaded file
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{register_number}_{timestamp}_{student_file.filename}"
        file_path = STUDENT_UPLOADS_DIR / filename
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(student_file.file, buffer)
            
        # 2. Extract text from PDF
        student_text = extract_text_from_pdf(str(file_path))
        if not student_text or len(student_text.strip()) < 50:
            raise HTTPException(status_code=400, detail="Could not extract enough text from the student paper. Ensure it is a valid PDF with selectable text.")
            
        # 3. Get official answer script
        cursor.execute(f"SELECT answer_data FROM answer_scripts WHERE question_paper_id = {placeholder} ORDER BY created_at DESC LIMIT 1", (paper_id,))
        script = cursor.fetchone()
        if not script:
            raise HTTPException(status_code=400, detail="No official answer script found for this paper. Generate it first.")
        
        answer_script = json.loads(script['answer_data'])
        
        # 4. AI Grading
        grading_result = grade_student_paper(answer_script, student_text)
        if not grading_result:
            raise HTTPException(status_code=500, detail="AI grading failed")
            
        marks_obtained = grading_result.get('total_marks_obtained', 0)
        total_marks = grading_result.get('total_max_marks', 100)
        result_status = "PASS" if marks_obtained >= (total_marks * 0.4) else "FAIL" # 40% Pass Mark
        
        # 5. Save evaluation to DB
        query = f"""
            INSERT INTO evaluations 
            (question_paper_id, student_name, register_number, department, marks_obtained, total_marks, result_status, evaluation_details, file_path)
            VALUES ({placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder}, {placeholder})
        """
        cursor.execute(query, (
            paper_id, student_name, register_number, department, 
            marks_obtained, total_marks, result_status, 
            json.dumps(grading_result), str(file_path)
        ))
        connection.commit()
        
        cursor.close()
        connection.close()
        return {"success": True, "result": result_status, "marks": marks_obtained}
        
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        traceback.print_exc()
        if connection: connection.close()
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/evaluations/results/{paper_id}", tags=["Evaluations"])
def get_results(paper_id: int, http_request: Request):
    """Get all evaluation results for a paper"""
    user = _current_user(http_request)
    connection = get_db_connection()
    if connection:
        _assert_paper_access(connection, user, paper_id)
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        cursor.execute(f"SELECT * FROM evaluations WHERE question_paper_id = {placeholder} ORDER BY created_at DESC", (paper_id,))
        results = cursor.fetchall()
        cursor.close()
        connection.close()
        return [dict(r) for r in results]
    except Exception as e:
        if connection: connection.close()
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/evaluations/report/{paper_id}", tags=["Evaluations"])
def get_report(paper_id: int, http_request: Request):
    """Get summary report for a paper"""
    user = _current_user(http_request)
    connection = get_db_connection()
    if connection:
        _assert_paper_access(connection, user, paper_id)
    try:
        cursor = get_cursor(connection)
        placeholder = get_placeholder()
        
        cursor.execute(f"SELECT * FROM evaluations WHERE question_paper_id = {placeholder}", (paper_id,))
        results = [dict(r) for r in cursor.fetchall()]
        
        if not results:
            return {"total_students": 0}
            
        total_students = len(results)
        pass_count = sum(1 for r in results if r['result_status'] == 'PASS')
        fail_count = total_students - pass_count
        average_marks = sum(r['marks_obtained'] for r in results) / total_students
        
        # Try to get answer script
        cursor.execute(f"SELECT answer_data FROM answer_scripts WHERE question_paper_id = {placeholder} ORDER BY created_at DESC LIMIT 1", (paper_id,))
        script = cursor.fetchone()
        answer_script = json.loads(script['answer_data']) if script else None

        cursor.close()
        connection.close()
        
        return {
            "total_students": total_students,
            "pass_count": pass_count,
            "fail_count": fail_count,
            "pass_percentage": (pass_count / total_students) * 100,
            "average_marks": average_marks,
            "results": results,
            "answer_script": answer_script
        }
    except HTTPException:
        raise
    except Exception as e:
        if connection: connection.close()
        raise HTTPException(status_code=500, detail=str(e))

if __name__ == "__main__":
    uvicorn.run(
        "main:app",
        host="0.0.0.0",
        port=int(os.getenv("PORT", "8010")),
        log_level="debug", # Add reload=True for hot-reload
        # log_level="info", in production mode
        access_log=True,
    )